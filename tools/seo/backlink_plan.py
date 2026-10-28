"""Backlink plan for swansura.com (user, 28 Oct 2026: 'Give me the worthy backlinks').
Builds docs/seo/Swansura_Backlink_Plan_2026-10-28.xlsx: prospects by tier, how to win each link, target page,
anchor idea, effort/cost, timing and a status column. Only white-hat routes (Google spam policy: no paid
dofollow links, PBNs, link exchanges or directory blasts). 'Checked 28 Oct' = site/page reached that day;
no authority scores are given because none were measured here (run Semrush Backlink Gap for numbers).
"""
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

S = 'https://www.swansura.com'
ROWS = [
 # priority, tier, target, url, why, how, link_to, anchor, effort, cost, timing, checked
 ('1', 'Brand partner', 'Deyga (official site)', 'https://deyga.in', 'The most relevant link possible: the brand itself naming Swansura as its UK partner. Also backs up the "Official UK partner" claim on the site.', 'Email your Deyga contact: ask for a "Buy in the UK" / international stockist line linking to the Deyga page on Swansura. No stockist page exists yet, so ask them to add one or a footer/FAQ line.', S + '/collections/milk-drop-brightening-serum-1', 'Deyga UK / Buy Deyga in the UK', 'Low', 'Free', 'Now', 'Site reached 28 Oct; no stockist page, no Swansura mention'),
 ('1', 'Brand partner', 'Vilvah (official site)', 'https://www.vilvah.com/pages/international-shipping', 'Same as Deyga: brand-to-stockist link, highly relevant, confirms the UK partner claim.', 'Ask Vilvah to add "UK customers: buy from our UK partner Swansura" on the International Shipping page (they ship abroad in USD, so a UK stockist line helps their UK customers too).', S + '/collections/milk-drop-brightening-serum', 'Vilvah UK', 'Low', 'Free', 'Now', 'Page reached 28 Oct; no Swansura mention'),
 ('1', 'Brand partner', 'anatomē stockists page', 'https://www.anatome.co/pages/stockists-1', 'anatomē has a real stockists page; Swansura sells its range (via Shopify Collective).', 'Ask anatomē (or via the Collective retailer contact) to list Swansura under UK online stockists.', S + '/collections/anatome', 'anatomē at Swansura', 'Low', 'Free', 'Now', 'Stockists page reached 28 Oct'),
 ('2', 'Brand partner', 'Highborn London', 'https://www.highbornlondon.com', 'UK brand, relevant skincare link.', 'Ask for a stockist mention (no stockists page found on the homepage); offer a short feature of their range on Swansura in return for the mention (no payment).', S + '/collections/highborn-london', 'Highborn London stockist', 'Low', 'Free', 'Now', 'Site reached 28 Oct; no stockists link found'),
 ('2', 'Brand partner', 'Jo Milano Paris / Game of Spades', 'https://jomilanoparis.com', 'Fragrance brand link; Swansura carries a large Jo Milano range.', 'Ask the UK distributor / brand contact whether they list online retailers; request a listing.', S + '/collections/luxury-collection-1', 'Jo Milano UK stockist', 'Low', 'Free', 'Now', 'Site reached 28 Oct; no stockists link found'),
 ('2', 'Awards', 'Free From Skincare Awards', 'https://freefromskincareawards.co.uk', 'UK awards covered by trade press; finalists/winners are listed with links. Entries usually open in January and close in February (2025: deadline 16 Feb).', 'Brands enter their products. Encourage Deyga/Vilvah to enter a UK-sold product and name Swansura as UK stockist; or check whether retailer entries are accepted for the 2027 cycle.', S + '/collections/milk-drop-brightening-serum', 'brand name', 'Medium', 'Entry fee (check)', 'Jan–Feb 2027', 'Site reached 28 Oct; 2027 dates not published yet'),
 ('2', 'Awards', 'The Beauty Shortlist Awards', 'https://www.thebeautyshortlist.com', 'Independent natural/organic beauty awards (UK-run, judges from UK/EU/US/AU). Strong fit for Indian natural skincare.', 'Same route as above: brand partners enter, Swansura named as UK stockist. Entries for 2026 ran 1–30 Sept 2025, so the 2027 window is likely Sept 2026 (missed) → plan for Sept 2027.', S + '/collections/natural-skincare', 'brand name', 'Medium', 'Entry fee (check)', 'Sept 2027', 'Site reached 28 Oct'),
 ('2', 'PR / journalists', '#journorequest on X', 'https://x.com/search?q=%23journorequest%20beauty', 'Free. Most #journorequests are UK-based; beauty and gift-guide requests are common in Nov–Dec.', 'Check daily for beauty, gifting, Indian/Ayurvedic skincare, perfume requests; reply fast with 2–3 lines, a product link and high-res images. Samples on request.', 'Product or Journal guide', 'product name', 'Medium', 'Free', 'Now (Christmas gift guides)', '—'),
 ('2', 'PR / journalists', 'ResponseSource (Journalist Enquiry Service)', 'https://www.responsesource.com', 'UK national/consumer press requests; Women\'s Interest & Beauty is its busiest category.', 'Subscribe to Women\'s Interest & Beauty; answer product/gift requests with samples.', 'Product pages', 'product name', 'Medium', 'Paid (check price)', 'Now', 'Site reached 28 Oct'),
 ('3', 'PR / journalists', 'Qwoted', 'https://www.qwoted.com', 'Journalist request platform; free tier.', 'Create a source profile (founder as UK natural-beauty retailer); pitch on beauty/retail/small-business requests.', 'Home / Journal', 'Swansura', 'Medium', 'Free tier', 'Now', 'Site reached 28 Oct'),
 ('3', 'PR / journalists', 'Featured.com', 'https://featured.com', 'Expert Q&A answers published on partner sites with a link.', 'Answer beauty/e-commerce questions as the founder; keep answers factual (no medical claims).', 'Home / Journal', 'Swansura', 'Medium', 'Free tier', 'Now', 'Site reached 28 Oct'),
 ('3', 'PR / journalists', 'SourceBottle', 'https://www.sourcebottle.com', 'Free call-outs incl. UK media and bloggers (gift guides, product reviews).', 'Sign up as a source; reply to UK beauty/gifting call-outs.', 'Product pages', 'product name', 'Low', 'Free', 'Now', 'Site reached 28 Oct'),
 ('2', 'Gift guides', 'UK beauty editors and bloggers – Christmas gift guides', '', 'Online gift guides run Nov–mid Dec and link straight to product pages.', 'Make a list of 20–30 UK beauty/lifestyle bloggers and online editors; send a short pitch with 5 giftable products (Christmas gift guide picks), prices, images, free UK delivery over £40; offer samples.', 'Product pages + /blogs/journal/christmas-gift-guide (once published)', 'product name', 'High', 'Samples + postage', 'Now – 10 Dec', '—'),
 ('2', 'Community media', 'Eastern Eye', 'https://www.easterneye.biz', 'British-Asian weekly; has covered South Asian beauty brands entering UK retail (e.g. Kay Beauty at Space NK).', 'Pitch a story: "the UK home of Indian natural skincare" – official UK partner of Deyga & Vilvah, ships from London; founder story; Christmas gifting angle.', 'Home or Deyga/Vilvah pages', 'Swansura', 'Medium', 'Free', 'Now', 'Site reached 28 Oct'),
 ('3', 'Community media', 'Asian Voice', 'https://www.asian-voice.com', 'British-Asian weekly, community and business stories.', 'Same story angle, local business focus (East London).', 'Home', 'Swansura', 'Medium', 'Free', 'Now', 'Site reached 28 Oct'),
 ('3', 'Community media', 'DESIblitz', 'https://www.desiblitz.com', 'British-Asian lifestyle site with beauty features.', 'Pitch an ingredient-led feature (amla, bhringraj, goat milk) using the Journal guides; offer quotes + images.', 'Journal guides', 'topic anchor', 'Medium', 'Free', 'After guides are published', 'Site reached 28 Oct'),
 ('3', 'Community media', 'Brown Beauty Talk', 'https://brownbeautytalk.com', 'Beauty site for brown skin; strong topical fit.', 'Pitch products for review/roundups (samples).', 'Product pages', 'product name', 'Medium', 'Samples', 'Now', 'Site reached 28 Oct'),
 ('2', 'Linkable content', 'Swansura Journal guides', S + '/blogs/journal', 'Guides on aloe vera gel, body scrub, body butter and hair oil target 5K–18K searches/month; useful pages earn links from bloggers and forums.', 'Publish the 5 drafts; then add an "Indian skincare ingredients A–Z" page built from the INCI data (unique on UK sites). Use them in every pitch above.', 'The guides', 'topic anchor', 'Medium', 'Free', 'This week', 'Drafts exist (unpublished)'),
 ('2', 'Creators', 'Shopify Collabs (installed)', 'https://apps.shopify.com/collabs', 'UK micro-creators post reviews on blogs/YouTube/TikTok; blog posts can carry links.', 'Invite 5–10 UK skincare creators with gifted products; ask for an honest review. Gifted/affiliate links must be disclosed and marked rel="sponsored" (Google + ASA/CAP rules).', 'Product pages', 'product name', 'Medium', 'Products', 'Nov–Dec', 'App installed'),
 ('3', 'Business profiles', 'Google Business Profile', 'https://business.google.com', 'Brand knowledge panel, Maps, reviews; not a backlink as such but a core trust signal.', 'Create/claim the profile (online retailer with a UK address; hide the address if customers don\'t visit); add website, categories, products.', 'Home', '—', 'Low', 'Free', 'Now', '—'),
 ('3', 'Business profiles', 'Bing Places', 'https://www.bing.com/forbusiness/', 'Bing/Copilot listings; can import from Google.', 'Import from Google Business Profile.', 'Home', '—', 'Low', 'Free', 'After GBP', 'Site reached 28 Oct'),
 ('3', 'Business profiles', 'Apple Business Connect', 'https://business.apple.com', 'Apple Maps / Siri listing.', 'Add the business with the same name, address, phone, website as everywhere else.', 'Home', '—', 'Low', 'Free', 'Now', 'Site reached 28 Oct'),
 ('3', 'Business profiles', 'Trustpilot (free business profile)', 'https://business.trustpilot.com', 'Profile page links to the site; reviews build trust for ads too.', 'Claim the free profile; invite customers after delivery.', 'Home', '—', 'Low', 'Free', 'Now', 'Site reached 28 Oct'),
 ('3', 'Business profiles', 'Judge.me public store page', 'https://judge.me', 'Judge.me (installed) can show a public store page with all reviews, linking back.', 'In the Judge.me app, turn on the public store page / Judge.me profile.', 'Home', '—', 'Low', 'Free', 'Now', 'App installed'),
 ('4', 'UK directories', 'Yell', 'https://www.yell.com', 'Largest UK business directory; citation + referral.', 'Free basic listing; keep name/address/phone identical everywhere.', 'Home', '—', 'Low', 'Free', 'Now', 'Real site (blocks bots)'),
 ('4', 'UK directories', 'Thomson Local', 'https://www.thomsonlocal.com', 'UK citation.', 'Free listing.', 'Home', '—', 'Low', 'Free', 'Now', 'Site reached 28 Oct'),
 ('4', 'UK directories', 'FreeIndex', 'https://www.freeindex.co.uk', 'UK citation with reviews.', 'Free listing.', 'Home', '—', 'Low', 'Free', 'Now', 'Real site (blocks bots)'),
 ('4', 'UK directories', 'Scoot', 'https://www.scoot.co.uk', 'UK citation (feeds other directories).', 'Free listing.', 'Home', '—', 'Low', 'Free', 'Now', 'Site reached 28 Oct'),
 ('4', 'UK directories', 'Cylex UK / Hotfrog', 'https://www.cylex-uk.co.uk', 'Extra citations; low value, do once and stop.', 'Free listings.', 'Home', '—', 'Low', 'Free', 'Later', 'Real sites (block bots)'),
 ('3', 'Local', 'Redbridge Council business support / local East London press', 'https://www.redbridge.gov.uk', 'Local authority and local news links (South Woodford, E18) are trusted.', 'Check the council\'s local business directory / support pages and local news titles for small-business features; pitch the founder story.', 'Home', 'Swansura', 'Medium', 'Free', 'Next month', 'Council site reached 28 Oct; directory not checked'),
 ('3', 'Research', 'Semrush Backlink Gap', 'https://www.semrush.com', 'Shows exactly which sites link to competitors but not to Swansura.', 'In your Semrush account: Backlink Gap for swansura.com vs skincupid.co.uk, indicrest.co.uk and 1–2 UK fragrance shops; export and add the relevant domains to this sheet.', '—', '—', 'Low', 'Your Semrush plan', 'This week', 'Skin Cupid site reached 28 Oct; Indicrest returned 500'),
]

AVOID = [
 ('Buying dofollow links / "guest posts" for a fee', 'Against Google\'s spam policy (link schemes); can lead to a manual penalty. Paid placements must be rel="sponsored".'),
 ('PBNs, link farms, "500 backlinks for £20" gigs', 'Spam; links get ignored at best, penalised at worst.'),
 ('Link exchanges ("link to me and I\'ll link to you") at scale', 'Link scheme under Google\'s policy.'),
 ('Mass directory / bookmarking / profile submissions', 'Low-value or spammy domains; only the handful of real UK directories in the plan are worth it.'),
 ('Blog comment / forum signature links', 'Nofollow and spammy; harms the brand.'),
 ('Undisclosed gifted or affiliate content', 'ASA/CAP rules: creator posts must be labelled (#ad / gifted); links rel="sponsored".'),
 ('Medical claims in pitches ("cures acne", "stops hair loss")', 'ASA rules; same no-claims line as the site copy.'),
]

PITCH = [
 ('Brand partner (Deyga / Vilvah)', 'Subject: UK stockist listing for Swansura\n\nHi <name>,\n\nAs your UK partner, we now carry <N> <brand> products at swansura.com, shipped from London with free UK delivery over £40. Many UK customers search for "<brand> UK", and a link from your site would help them find an official UK seller.\n\nCould you add Swansura to your international/stockists page, or a line such as "UK customers: shop <brand> at Swansura" linking to <url>?\n\nThank you,\n<name>, Swansura – The Body Studio'),
 ('Gift guide (editor / blogger)', 'Subject: Christmas gift idea: Indian natural skincare from a UK studio\n\nHi <name>,\n\nI loved your <guide/article>. For this year\'s gift guides, Swansura is the UK home of Deyga and Vilvah natural skincare, alongside niche perfume. Five picks under £50: <product – price – link> ×5.\n\nAll ship from London, free UK delivery over £40. Hi-res images attached; happy to send samples.\n\nBest,\n<name>'),
 ('Community media (Eastern Eye / Asian Voice)', 'Subject: East London studio brings Indian natural skincare to the UK\n\nHi <name>,\n\nSwansura – The Body Studio, based in South Woodford, is the official UK partner of Indian natural skincare brands Deyga and Vilvah, alongside 60+ perfume houses. <Founder> started it to <reason>. Happy to share the story, images and products for review.\n\nBest,\n<name>'),
]

wb = Workbook()
ws = wb.active
ws.title = 'Backlink plan'
hdr = ['Priority', 'Type', 'Target', 'URL', 'Why it is worth it', 'How to get the link', 'Swansura page to link', 'Anchor idea', 'Effort', 'Cost', 'Timing', 'Checked', 'Status', 'Link live (URL)', 'Notes']
ws.append(hdr)
for r in sorted(ROWS, key=lambda r: r[0]):
    ws.append(list(r) + ['To do', '', ''])
head = PatternFill('solid', fgColor='1B4D3E')
for c in ws[1]:
    c.font = Font(bold=True, color='FFFFFF'); c.fill = head; c.alignment = Alignment(wrap_text=True, vertical='top')
widths = [8, 16, 30, 38, 48, 58, 40, 24, 9, 16, 18, 26, 12, 30, 30]
for i, w in enumerate(widths):
    ws.column_dimensions[chr(65 + i)].width = w
for row in ws.iter_rows(min_row=2):
    for c in row:
        c.alignment = Alignment(wrap_text=True, vertical='top')
ws.freeze_panes = 'D2'
ws.auto_filter.ref = ws.dimensions
dv = DataValidation(type='list', formula1='"To do,Contacted,Follow-up sent,Won,Declined,Not relevant"', allow_blank=True)
ws.add_data_validation(dv); dv.add(f'M2:M{ws.max_row}')
tints = {'1': 'E4EEE9', '2': 'F4EADA', '3': 'FFFFFF', '4': 'F5F2ED'}
for row in ws.iter_rows(min_row=2):
    for c in row:
        c.fill = PatternFill('solid', fgColor=tints[row[0].value])

ws2 = wb.create_sheet('Pitch templates')
ws2.append(['For', 'Email template (fill the <…> parts)'])
for r in PITCH:
    ws2.append(list(r))
ws2.column_dimensions['A'].width = 30; ws2.column_dimensions['B'].width = 110
for row in ws2.iter_rows():
    for c in row:
        c.alignment = Alignment(wrap_text=True, vertical='top')
for c in ws2[1]:
    c.font = Font(bold=True, color='FFFFFF'); c.fill = head

ws3 = wb.create_sheet('Avoid')
ws3.append(['Do not do', 'Why'])
for r in AVOID:
    ws3.append(list(r))
ws3.column_dimensions['A'].width = 55; ws3.column_dimensions['B'].width = 90
for row in ws3.iter_rows():
    for c in row:
        c.alignment = Alignment(wrap_text=True, vertical='top')
for c in ws3[1]:
    c.font = Font(bold=True, color='FFFFFF'); c.fill = head

out = 'docs/seo/Swansura_Backlink_Plan_2026-10-28.xlsx'
wb.save(out)
print(out, len(ROWS), 'prospects')
