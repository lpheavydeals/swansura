"""Old draft products -> client decision sheet (31 Oct 2026).

Input : docs/data/old_drafts_2026-10-31.json  (Admin GraphQL: status:draft AND -tag:needs-price)
        live storefront /products.json (active products, fetched here)
Output: docs/Swansura_Old_Draft_Products_2026-10-31.xlsx

Each draft gets a suggested action + reason and a Decision dropdown for the client.
Run: python3 tools/shopify/old_drafts_sheet.py
"""
import json, re, collections, urllib.request
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

SRC = 'docs/data/old_drafts_2026-10-31.json'
OUT = 'docs/Swansura_Old_Draft_Products_2026-10-31.xlsx'
ADMIN = 'https://admin.shopify.com/store/d2rqgy-5e/products/'

drafts = json.load(open(SRC))

act = []
for p in (1, 2, 3, 4):
    req = urllib.request.Request(f'https://www.swansura.com/products.json?limit=250&page={p}',
                                 headers={'User-Agent': 'Mozilla/5.0 Chrome/130'})
    batch = json.load(urllib.request.urlopen(req))['products']
    if not batch:
        break
    act += batch
ACT = {a['handle']: a for a in act}

# ---- duplicates among the drafts -------------------------------------------------
STOP = set('''riiffs riiff parfums parfum perfume perfumes nusuk rihanah eau de du edp edt toilette
spray natural for men women man woman unisex ladies homme femme pour long lasting longlasting luxury
fragrance fragrances arabian scent the and with 100ml 100 ml 3 4 oz by exclusive aromix masculine'''.split())
SYN = {'bleu': 'blue', 'absolue': 'absolu', 'golden': 'gold'}

def key(t):
    t = t.lower().replace('’', "'")
    giftset = 'gift set' in t
    t = re.split(r'\s[|–]\s', t)[0]
    w = [SYN.get(x, x) for x in re.findall(r'[a-z0-9]+', t)]
    w = tuple(x for x in w if x not in STOP)
    return w + (('GIFTSET',) if giftset else ())

def gender(t):
    t = t.lower()
    if 'unisex' in t or ('men and women' in t) or ('men & women' in t):
        return 'u'
    if re.search(r'\bwomen\b|\bfemme\b|\bladies\b', t):
        return 'w'
    if re.search(r'\bmen\b|\bman\b|\bhomme\b', t):
        return 'm'
    return ''

groups = collections.defaultdict(list)
for p in drafts:
    groups[key(p['title'])].append(p)

dup_of = {}   # id -> title of the listing to keep
for k, g in groups.items():
    if len(g) < 2:
        continue
    gs = {gender(p['title']) for p in g} - {'', 'u'}
    if gs == {'m', 'w'} and len(g) == 2:   # e.g. Wonder of You for Men / for Women: two scents
        continue
    keep = max(g, key=lambda p: (p['mediaCount']['count'], len(p['title'])))
    for p in g:
        if p is not keep:
            dup_of[p['id']] = keep['title']

# ---- drafts that already exist as live products -----------------------------------
LIVE = {   # draft title start -> live handle (checked by hand against the storefront)
    '10% Vitamin C Face Serum 20ml': None, 'Beautifying Serum 30 ml': None,
    'Charcoal Bath Bar 130 g': None, "Deyga's Rose & Mulethi": None, "Deyga's Under Eye Cream": None,
    'Acqua di Parma Holiday Discovery Set': None,
    'Giorgio Armani Code EDT 125ml': 'giorgio-armani-code-eau-de-toilette-125ml-refillable-men-s-woody-aromatic-fragrance',
    'Game Of Spades Moon Unisex': 'game-of-spades-moon-p-100ml-unisex',
    'Game Of Spades Blind-Bid': 'jo-milano-blind-bid-p-100-ml-unisex',
}

def live_match(t):
    for start, h in LIVE.items():
        if t.startswith(start):
            if h:
                return ACT[h]['title'] if h in ACT else h
            # find by core words
            core = re.split(r'\s[|–]\s', t)[0].replace("Deyga's", '').strip().lower()
            for a in act:
                if core[:22] in a['title'].lower():
                    return a['title']
    return None

def brand_of(p):
    t, v = p['title'], p['vendor']
    if re.search(r'riiffs|nusuk|rihanah', t + ' ' + v, re.I):
        return 'RiiFFS / Nusuk'
    if t.lower().startswith('game of spades'):
        return 'Game of Spades (Jo Milano)'
    if 'maison de milan' in t.lower():
        return 'Maison de Milan'
    if 'deyga' in t.lower():
        return 'Deyga'
    if 'vilvah' in t.lower():
        return 'Vilvah'
    if v == 'Giorgia Armani':
        return 'Giorgio Armani'
    return v

def suggest(p):
    t, stock = p['title'], p['totalInventory']
    b = brand_of(p)
    m = live_match(t)
    if m:
        return 'Delete', f'Already live as "{m}". This draft is a copy.'
    if p['id'] in dup_of and ('femme' in t.lower()) != ('femme' in dup_of[p['id']].lower()):
        return 'Check', f'Maybe the same perfume as draft "{dup_of[p["id"]]}". Keep one if they are the same.'
    if p['id'] in dup_of:
        return 'Delete', f'Duplicate draft of "{dup_of[p["id"]]}" (same perfume listed twice).'
    if t.startswith('Juicy Couture Viva La Juicy'):
        return 'Delete', ('Already live: the Viva La Juicy Rosé listing has a 50ml size (£31.99). '
                          f'Move this draft\'s stock ({stock}) to that size, then delete.')
    if t.startswith('Deyga Natural Hair Care Kit'):
        return 'Check', ('Live has "Deyga Hair Care Kit" (50 / 220 sizes, £18.99 / £36.99). This draft is £24.99 '
                         f'with {stock} in stock: same kit (delete) or a different set (publish)?')
    if b == 'Game of Spades (Jo Milano)':
        if t.startswith('Game Of Spades Rouge Extrait'):
            return 'Check', ('Live has "Game of Spades Rouge Unisex Eau de Parfum". Publish only if '
                             f'Rouge Extrait is a different perfume (stock {stock}); vendor -> Jo Milano.')
        if stock > 0:
            return 'Finish & publish', (f'Not on the site yet and {stock} in stock. Needs vendor Jo Milano, '
                                        'short title, SEO title/description.')
        return 'Keep as draft', 'Not on the site, 0 in stock. Publish when stock arrives.'
    if b in ('Deyga', 'Vilvah'):
        if 'Aloe Vera Gel Pack of 2' in t:
            return 'Finish & publish', f'{stock} in stock. Vendor -> Vilvah, short title, SEO.'
        if 'Melt-in-Milk' in t:
            return 'Check', ('0 in stock. Vilvah now sells Skin Finish Sunscreen (live). '
                             'Delete unless you will restock Melt-in-Milk.')
        if 'Hair Strengthening Kit' in t:
            return 'Check', '0 in stock, 4 photos, vendor Generic. Publish only if Deyga still makes this kit.'
    if b == 'RiiFFS / Nusuk':
        return 'Keep as draft', ('0 in stock. Never sold on swansura.com. Delete all RiiFFS/Nusuk drafts '
                                 'if you will not stock the brand.')
    if b == 'Chris Adams':
        return 'Keep as draft', '0 in stock. Delete if you will not stock Chris Adams.'
    if b == 'Maison de Milan':
        return 'Keep as draft', '0 in stock, 1 photo, vendor Generic. Delete if you will not stock Maison de Milan.'
    return 'Check', ''

rows = []
for p in drafts:
    act_, why = suggest(p)
    rows.append({
        'brand': brand_of(p), 'title': p['title'], 'vendor': p['vendor'],
        'price': float(p['priceRangeV2']['minVariantPrice']['amount']),
        'stock': p['totalInventory'], 'images': p['mediaCount']['count'],
        'variants': p['variantsCount']['count'],
        'created': p['createdAt'][:10], 'updated': p['updatedAt'][:10],
        'action': act_, 'why': why, 'link': ADMIN + p['id'].split('/')[-1],
    })
ORDER = ['Deyga', 'Vilvah', 'Game of Spades (Jo Milano)', 'Juicy Couture', 'Acqua di Parma',
         'Giorgio Armani', 'Chris Adams', 'Maison de Milan', 'RiiFFS / Nusuk']
rows.sort(key=lambda r: (ORDER.index(r['brand']) if r['brand'] in ORDER else 99, r['title'].lower()))

# ---- workbook ---------------------------------------------------------------------
EM, GOLD, LINE = '1B4D3E', 'B5843A', 'E3DACD'
FILL = {'Delete': 'F3E6E0', 'Finish & publish': 'E4EEE9', 'Check': 'F4EADA', 'Keep as draft': 'FFFFFF'}
thin = Side(style='thin', color=LINE)
B = Border(top=thin, bottom=thin, left=thin, right=thin)
F = 'Plus Jakarta Sans'

wb = Workbook()
s = wb.active
s.title = 'Summary'
s['A1'] = 'Swansura – old draft products to review'
s['A1'].font = Font(name=F, size=16, bold=True, color=EM)
s['A2'] = ('157 draft products in Shopify that are not on the website (created Oct 2025, before the new site). '
           'The 10 drafts waiting for prices (tag needs-price) are not in this list. '
           'Pick a decision for each row on the "Drafts" tab; nothing changes in the store until you confirm.')
s['A2'].alignment = Alignment(wrap_text=True, vertical='top')
s.merge_cells('A2:D2')
s.row_dimensions[2].height = 62
s['A4'], s['B4'], s['C4'], s['D4'] = 'Suggested action', 'Products', 'In stock (units)', 'What it means'
MEAN = {'Delete': 'Copy of a live product or of another draft. Safe to delete.',
        'Finish & publish': 'Has stock and is not on the site. We tidy vendor, title and SEO, then publish.',
        'Check': 'Needs your answer first (see the reason column).',
        'Keep as draft': 'No stock and never sold here. Keep hidden, or delete if you will not stock the brand.'}
r = 5
for a in ['Delete', 'Finish & publish', 'Check', 'Keep as draft']:
    sel = [x for x in rows if x['action'] == a]
    s.cell(r, 1, a); s.cell(r, 2, len(sel)); s.cell(r, 3, sum(max(x['stock'], 0) for x in sel)); s.cell(r, 4, MEAN[a])
    for c in range(1, 5):
        s.cell(r, c).fill = PatternFill('solid', fgColor=FILL[a]); s.cell(r, c).border = B
        s.cell(r, c).alignment = Alignment(wrap_text=True, vertical='top')
    r += 1
s.cell(r, 1, 'Total').font = Font(name=F, bold=True)
s.cell(r, 2, len(rows)).font = Font(name=F, bold=True)
s.cell(r, 3, sum(max(x['stock'], 0) for x in rows)).font = Font(name=F, bold=True)
r += 2
s.cell(r, 1, 'By brand').font = Font(name=F, size=12, bold=True, color=EM); r += 1
for c, h in enumerate(['Brand', 'Drafts', 'In stock (units)', 'Suggested'], 1):
    s.cell(r, c, h)
hdr_rows = [4, r]
r += 1
for b in ORDER:
    sel = [x for x in rows if x['brand'] == b]
    if not sel:
        continue
    cnt = collections.Counter(x['action'] for x in sel)
    s.cell(r, 1, b); s.cell(r, 2, len(sel)); s.cell(r, 3, sum(max(x['stock'], 0) for x in sel))
    s.cell(r, 4, ', '.join(f'{v} {k.lower()}' for k, v in cnt.most_common()))
    for c in range(1, 5):
        s.cell(r, c).border = B
    r += 1
for hr in hdr_rows:
    for c in range(1, 5):
        x = s.cell(hr, c)
        x.font = Font(name=F, bold=True, color='FFFFFF'); x.fill = PatternFill('solid', fgColor=EM); x.border = B
for w, col in zip([24, 10, 16, 80], 'ABCD'):
    s.column_dimensions[col].width = w

d = wb.create_sheet('Drafts')
H = ['Brand', 'Product title (as in Shopify)', 'Vendor in Shopify', 'Price £', 'Stock', 'Photos',
     'Sizes', 'Created', 'Last edited', 'Suggested action', 'Why', 'Your decision', 'Notes', 'Open in Shopify']
W = [20, 58, 16, 9, 7, 7, 6, 11, 11, 16, 62, 18, 28, 16]
for c, (h, w) in enumerate(zip(H, W), 1):
    x = d.cell(1, c, h)
    x.font = Font(name=F, bold=True, color='FFFFFF'); x.fill = PatternFill('solid', fgColor=EM)
    x.alignment = Alignment(wrap_text=True, vertical='center'); x.border = B
    d.column_dimensions[get_column_letter(c)].width = w
d.row_dimensions[1].height = 30
for i, x in enumerate(rows, 2):
    vals = [x['brand'], x['title'], x['vendor'], x['price'], x['stock'], x['images'], x['variants'],
            x['created'], x['updated'], x['action'], x['why'], '', '', 'Open']
    for c, v in enumerate(vals, 1):
        cell = d.cell(i, c, v)
        cell.font = Font(name=F, size=10)
        cell.alignment = Alignment(wrap_text=c in (2, 11, 13), vertical='top')
        cell.border = B
    d.cell(i, 4).number_format = '£#,##0.00'
    d.cell(i, 10).fill = PatternFill('solid', fgColor=FILL[x['action']])
    d.cell(i, 12).fill = PatternFill('solid', fgColor='FFF8E8')
    lk = d.cell(i, 14); lk.hyperlink = x['link']; lk.font = Font(name=F, size=10, color=EM, underline='single')
dv = DataValidation(type='list', formula1='"Delete,Finish & publish,Keep as draft,Not sure"', allow_blank=True)
d.add_data_validation(dv)
dv.add(f'L2:L{len(rows) + 1}')
d.freeze_panes = 'C2'
d.auto_filter.ref = f'A1:N{len(rows) + 1}'

wb.save(OUT)
print(OUT, len(rows), collections.Counter(x['action'] for x in rows))
