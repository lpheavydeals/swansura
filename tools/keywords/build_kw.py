import openpyxl, re, json, sys
SRC, OUT = sys.argv[1], sys.argv[2]
wb = openpyxl.load_workbook(SRC, data_only=True)
BRANDFIX = {'tomford':'Tom Ford','hugoboss':'Hugo Boss','abercrombie& fitch':'Abercrombie & Fitch','phillip plein':'Philipp Plein','calvin klein':'Calvin Klein',
 'dior':'Dior','antonio':'Antonio Banderas','chloe':'Chloé','sap':'SAP','dkny':'DKNY','anatomē':'anatomē','la maison':'La Maison'}
NOBRAND = {'generic','swansura - the body studio'}
def canon(b): b=(b or '').strip(); return BRANDFIX.get(b.lower(), b)
import unicodedata
def norm(s): return re.sub(r'\s+',' ', s.replace('’',"'").replace('&amp;','&')).strip()
def plain(s): return ''.join(c for c in unicodedata.normalize('NFD', s) if unicodedata.category(c)!='Mn')
SIZE = r'\b\d+(\.\d+)?\s*(ml|g|gm|gms|grams?|oz)\b|\b\d+\s*x\s*\d+\s*(ml|g)\b'
CONC = r"\b(extrait de parfum|eau de parfum|eau de toilette|eau de cologne|eau fraiche|edp|edt|edc)\b"
def frag_core(name, brand):
    n = norm(name)
    n = re.split(r'\s[–—|]\s|\s[–—]|–', n)[0]            # drop descriptor tails
    n = re.sub(r'\s-\s.*$', '', n)                        # "Pour Femme - Eau de Parfum for Women"
    n = re.sub(r'\+.*$', '', n)                           # "+15ml travel spray"
    n = re.sub(SIZE, ' ', n, flags=re.I)
    conc = re.search(CONC, n, flags=re.I)
    n = re.sub(CONC, ' ', n, flags=re.I)
    n = re.sub(r"\b(for (women|men|him|her)|unisex|perfume|fragrance|spray|luxury|long lasting|women's|men's|womens|mens)\b", ' ', n, flags=re.I)
    n = re.sub(r'\bparfum\b\s*$', '', n.strip(), flags=re.I)
    n = re.sub(r'[()]', ' ', n); n = norm(n).strip(' -,')
    b = canon(brand)
    if b.lower() not in NOBRAND and not n.lower().startswith(b.lower().split()[0].lower()):
        n = f'{b} {n}'
    return plain(n).lower(), (conc.group(1).lower() if conc else '')
def skin_core(name, brand):
    n = norm(name); n = re.sub(SIZE, ' ', n, flags=re.I); n = re.sub(r'\b\d+\s*(g|ml)\b', ' ', n, flags=re.I)
    n = re.sub(r"^(vilvah|deyga'?s?)\s+", '', n.strip(), flags=re.I); n = norm(n)
    return f'{brand} {n}'.lower()
CONCERN = {'Anti-Ageing':'anti ageing','Anti-Dandruff':'dandruff','Dark Circles':'dark circles','Dry Lips':'dry lips','Dry Skin':'dry skin','Dull Skin':'dull skin','Frizzy Hair':'frizzy hair','Hair Thinning & Loss':'hair loss','Oily Skin':'oily skin','Acne & Blemishes':'acne','Sun Protection':None,'Damaged Skin Barrier':'damaged skin barrier','Lip Exfoliation':None}
TYPE = {'Bath & Body Wash':'body wash','Body Butters':'body butter','Body Cream & Lotions':'body lotion','Body Oil':'body oil','Body Scrub & Exfoliators':'body scrub','Cleansers':'face wash','Conditioner':'conditioner','Deodorant':'natural deodorant','Face Masks':'face mask','Gift Sets & Combos':'skincare gift set','Hair Masks':'hair mask','Hair Oil':'hair oil','Hair Serums':'hair serum','Lip Care':'lip balm','Moisturisers':'moisturiser','SPF for Face':'sunscreen','Serums':'face serum','Shampoo':'shampoo','Toners':'toner','Skincare Treatments':'skincare','Body Care':'body care','Hair Treatments':'hair treatment'}
prods, kws = [], {}
def add(k, kind, tier, pid=None):
    k = plain(norm(k)).lower().strip()
    if len(k) < 3: return
    e = kws.setdefault(k, {'kw':k,'kind':kind,'tier':tier,'pids':[]})
    e['tier'] = min(e['tier'], tier)
    if pid is not None and pid not in e['pids']: e['pids'].append(pid)
ws = wb['Vilvah & Deyga']; hdr = [c.value for c in ws[1]]
for r in ws.iter_rows(min_row=2, values_only=True):
    if r[0] is None or not r[5]: continue
    d = dict(zip(hdr, r)); pid = f"VD-{int(d['S.No']):03d}"; b = d['Brand'].strip()
    core = skin_core(d['Product Name'], b); typ = TYPE.get(d['Product Type'], (d['Product Type'] or '').lower()); con = CONCERN.get(d['Skin / Hair Concern'], (d['Skin / Hair Concern'] or '').lower() or None)
    p = dict(id=pid, sheet='Vilvah & Deyga', row=d['S.No'], brand=b, brand_fix='', cat=d['Category'], type=d['Product Type'], concern=d['Skin / Hair Concern'], name=norm(d['Product Name']), size=d['Size'], price=d['MRP (£)'], core=core,
             kws=[core, f'{core} uk', f'buy {core} online', f'{core} price'], cluster=[f'{b.lower()} {typ}'] + ([f'{typ} for {con}'] if con and typ and 'gift' not in typ else []) + ([f'natural {typ} uk'] if typ and 'gift' not in typ else []))
    prods.append(p)
    add(core,'Product',1,pid); add(f'{core} uk','Product + UK',2,pid); add(f'buy {core} online','Buy online',3,pid); add(f'{core} price','Price',3,pid)
    for c in p['cluster']: add(c,'Category / concern',1,pid)
for b in ['vilvah','deyga']:
    for k in [b, f'{b} uk', f'{b} products', f'{b} skincare', f'{b} products uk', f'buy {b} products online', f'{b} official website']: add(k,'Brand',1)
ws = wb['Fragrance Catalogue']; hdr = [c.value for c in ws[1]]
for r in ws.iter_rows(min_row=2, values_only=True):
    if r[0] is None or not r[6]: continue
    d = dict(zip(hdr, r)); pid = f"FR-{int(d['S.No']):03d}"; braw = (d['Brand'] or '').strip(); b = canon(braw)
    core, conc = frag_core(d['Product Name'], braw)
    if not conc: conc = {'Eau de Parfum (EDP)':'eau de parfum','Eau de Toilette (EDT)':'eau de toilette','Extrait de Parfum':'extrait de parfum','Parfum':'parfum'}.get(d['Fragrance Type'],'')
    gender = {"Women's Fragrance":'for women',"Men's Fragrance":'for men','Unisex Fragrance':'unisex'}.get(d['Subcategory'],'')
    nm = norm(d['Product Name']); first = nm.lower().split()[0] if nm else ''
    note = []
    if braw.lower() in NOBRAND: note.append('Real brand missing')
    elif b.lower().split()[0] not in nm.lower() and not core.startswith(b.lower()): note.append('Check brand')
    elif b.lower() != braw.lower(): note.append(f'Spelling: {b}')
    known = ['azzaro','gucci','dior','chanel','armani','versace','paco','lancome','ysl','prada','givenchy','hugo','boss','calvin','bvlgari','burberry']
    if braw.lower() not in NOBRAND and first in known and first not in b.lower(): note = [f'Name says {first.title()}, brand says {braw}']
    kl = [core] + ([f'{core} {conc}'] if conc and 'set' not in core else []) + [f'{core} uk', f'buy {core} online', f'{core} price']
    fam = (d['Scent Family'] or '').lower()
    cl = []
    if braw.lower() not in NOBRAND:
        cl = [f'{b.lower()} perfume'] + ([f'{b.lower()} perfume {gender}'] if gender in ('for women','for men') else [])
    if fam: cl.append(f'{fam} perfume' + (f' {gender}' if gender in ('for women','for men') else '') + ' uk')
    p = dict(id=pid, sheet='Fragrance Catalogue', row=d['S.No'], brand=braw, brand_fix='; '.join(note), cat=d['Category'], type=d['Fragrance Type'], concern=f"{d['Scent Family']} · {d['Subcategory']}", name=nm, size=d['Size'], price=d['MRP (£)'], core=core, kws=kl, cluster=cl)
    prods.append(p)
    add(core,'Product',1,pid)
    for k in kl[1:]: add(k, 'Product + type' if k.endswith(conc) and conc else 'Product + UK' if k.endswith(' uk') else 'Buy online' if k.startswith('buy ') else 'Price', 2 if (k.endswith(' uk') or (conc and k.endswith(conc))) else 3, pid)
    for c in cl: add(c,'Brand / family',1,pid)
json.dump({'products':prods,'keywords':list(kws.values())}, open(OUT,'w'), indent=1, ensure_ascii=False)
from collections import Counter
print('products',len(prods),'keywords',len(kws), Counter(k['tier'] for k in kws.values()))
