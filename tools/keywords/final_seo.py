"""'Final SEO' tab: one row per product with the final H1, meta title and meta description (fragrance = data-driven pass; Vilvah/Deyga = VD Content)."""
import openpyxl, sys, re, unicodedata
from collections import defaultdict
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import FormulaRule
WB = sys.argv[1]
v = openpyxl.load_workbook(WB, data_only=True)
rows = lambda ws: [dict(zip([c.value for c in ws[1]], r)) for r in ws.iter_rows(min_row=2, values_only=True)]
P, K, VDC = rows(v['Products']), rows(v['Keyword List']), {r['ID']: r for r in rows(v['VD Content'])}
VOL = {r[0]: (r[2] if isinstance(r[2], (int, float)) else 0) for r in v['Semrush Data'].iter_rows(min_row=2, values_only=True) if r[0]}
plain = lambda s: ''.join(c for c in unicodedata.normalize('NFD', s) if unicodedata.category(c) != 'Mn').lower()
by_pid = defaultdict(list)
for k in K:
    for pid in str(k['Product IDs'] or '').replace(' …', '').split(', '):
        if pid: by_pid[pid].append(k['Keyword'])
CONC = {'Eau de Parfum (EDP)': 'EDP', 'Eau de Toilette (EDT)': 'EDT', 'Extrait de Parfum': 'EXTRAIT', 'Parfum': 'PARFUM', 'Parfum / Extrait de Parfum': 'PARFUM', 'Eau de Cologne (EDC)': 'EDC'}
CONC_WORD = {'EDP': 'eau de parfum', 'EDT': 'eau de toilette', 'EXTRAIT': 'extrait de parfum', 'PARFUM': 'parfum', 'EDC': 'eau de cologne'}
def conc_of(k):
    k = ' ' + k + ' '
    if 'eau de parfum' in k or ' edp ' in k: return 'EDP'
    if 'eau de toilette' in k or ' edt ' in k: return 'EDT'
    if 'extrait' in k: return 'EXTRAIT'
    if 'cologne' in k: return 'EDC'
    if ' parfum ' in k: return 'PARFUM'
    return None
def gender_of(p):
    c = p['Concern / scent'] or ''
    return 'women' if "Women" in c else 'men' if "Men's" in c else 'unisex'
BAD = re.compile(r'review|dupe|notes|price|sale|tester|similar|\bvs\b|sample|20\d\d|fragrantica|attar|amazon|boots|superdrug|cheap')
def secondary(p):
    core, pc, g = p['Core keyword'], CONC.get(p['Type']), gender_of(p)
    best = (None, 0)
    for k in by_pid.get(p['ID'], []):
        if k == core or BAD.search(k) or k.startswith('buy ') or not k.startswith(core + ' ') or k.endswith(' uk'): continue
        kc = conc_of(k)
        if kc and kc != pc: continue
        if ('for women' in k or 'pour femme' in k) and g == 'men': continue
        if ('for men' in k or 'pour homme' in k) and g == 'women': continue
        if VOL.get(k, 0) > best[1]: best = (k, VOL[k])
    return best[0] or ''
groups = defaultdict(list)
for p in P:
    if p['Source sheet'] != 'Vilvah & Deyga': groups[p['Core keyword']].append(p)
SMALL = {'and', 'or', 'of', 'de', 'di', 'du', 'la', 'le', 'for', 'the', 'with', 'in', 'a', '&', 'pour', 'des', 'da'}
def fit(cands, lo, hi, target):
    cands = [re.sub(r'\s+', ' ', c).strip() for c in cands if c]
    ok = [c for c in cands if lo <= len(c) <= hi]
    if ok: return ok[0] if target == 'first' else max(ok, key=len)
    under = [c for c in cands if len(c) <= hi]
    return max(under, key=len) if under else cands[-1][:hi]
BRANDFIX = {'tomford':'Tom Ford','hugoboss':'Hugo Boss','abercrombie& fitch':'Abercrombie & Fitch','phillip plein':'Philipp Plein','calvin klein':'Calvin Klein','dior':'Dior','antonio':'Antonio Banderas','chloe':'Chloé','sap':'SAP','dkny':'DKNY','la maison':'La Maison'}
NOBRAND = {'generic', 'swansura - the body studio'}
sig = defaultdict(list)
for p in P: sig[(p['Product name'].split(' – ')[0].strip().lower()[:60], str(p['Size']).replace(' ', '').lower())].append(p)
def fix_brand(h1, braw):
    b = BRANDFIX.get(braw.strip().lower(), braw.strip())
    if braw.strip().lower() in NOBRAND: return h1
    first = plain(b).split()[0]
    return h1 if first in plain(h1) else f'{b} {h1}'
out = []
for p in P:
    if p['Source sheet'] == 'Vilvah & Deyga':
        r = VDC[p['ID']]; h1v = r['H1']; notev = 'Target keyword is shared with its collection page (' + str(r['Collection page (link to)']) + ')'
        same = [x for x in P if x['Source sheet'] == 'Vilvah & Deyga' and VDC[x['ID']]['H1'] == h1v]
        szv = str(p['Size'] or '').replace(' ', '')
        if len(same) > 1:
            if szv and szv.lower() not in ('none', 'notspecified') and len({str(x['Size']) for x in same}) == len(same): h1v = f'{h1v} {szv}'
            else: notev += '; Duplicate listing (same name/size): check the sheet'
            notev += '; Size variants: merge as variants of one product'
        kw = r['Target keyword']; KWt = ' '.join(w if w in ('for', 'and', 'of') else w[:1].upper() + w[1:] for w in kw.split())
        redundant = set(kw.split()) <= set(plain(h1v).replace('-', ' ').split()) | {'natural', 'for'}
        tv = [f'{h1v} | {KWt}' if not redundant else '', f'{h1v} – {KWt} | Swansura' if not redundant else '', f'{h1v} | Buy Online UK | Swansura', f"{h1v} | {p['Brand']} UK | Swansura", f'{h1v} | Swansura UK']
        title_v = r['Meta title'] if (h1v == r['H1'] and not redundant) else fit(tv, 50, 60, 'first')
        out.append(dict(p=p, primary=kw, secondary=p['Brand'].lower(), h1=h1v, title=title_v, desc=r['Meta description'], note=notev))
        continue
    g, pc = gender_of(p), CONC.get(p['Type'])
    h1 = fix_brand(p['Draft H1'], p['Brand']); size = str(p['Size'] or '').replace(' ', '')
    if size.lower() in ('none', 'notspecified', 'n/a', 'na', '-'): size = ''
    grp = groups[p['Core keyword']]; note = ''
    if len(grp) > 1:
        genders = {gender_of(x) for x in grp}
        if len(genders) > 1 and g in ('women', 'men') and not re.search(r'for (women|men|him|her)|pour (homme|femme)', h1, re.I):
            h1 = f"{h1} for {'Women' if g == 'women' else 'Men'}"
        elif size and size.lower() not in h1.lower():
            h1 = f'{h1} {size}'
        note = f'Shares the keyword "{p["Core keyword"]}" with {len(grp) - 1} other listing(s): consider merging sizes as variants of one product'
    sec = secondary(p)
    twins = sig[(p['Product name'].split(' – ')[0].strip().lower()[:60], str(p['Size']).replace(' ', '').lower())]
    if len(twins) > 1 and len({x['MRP (£)'] for x in twins}) > 1:
        note = (note + '; ' if note else '') + 'Same name and size listed at different prices (' + ', '.join(f"£{x['MRP (£)']}" for x in twins) + '): check size/price in the sheet'
    h1b = re.sub(r'\s*\b(Extrait de Parfum|Eau de Parfum|Eau de Toilette|Eau de Cologne|EDP|EDT|Parfum)\b\s*', ' ', h1, flags=re.I).strip()
    cw = {'EDP': 'EDP', 'EDT': 'EDT', 'EXTRAIT': 'Extrait', 'PARFUM': 'Parfum', 'EDC': 'Cologne'}.get(pc, '')
    sz = '' if size and size.lower() in h1.lower() else size
    title = fit([f'{h1} {sz} | Buy Online UK | Swansura' if sz else '', f'{h1} | Buy Online UK | Swansura', f'{h1} {sz} | Swansura UK' if sz else '', f'{h1b} {cw} {sz} | Swansura UK' if cw else '', f'{h1} | Swansura UK', f'{h1b} | Swansura UK'], 50, 60, 'first')
    price = p['MRP (£)']; pr = f' for £{price:.2f}' if isinstance(price, (int, float)) else ''
    fam = (p['Concern / scent'] or '').split(' · ')[0].lower()
    gw = {'women': "women's", 'men': "men's", 'unisex': 'unisex'}[g]
    kind = CONC_WORD.get(pc, 'fragrance') if 'set' not in p['Type'].lower() else 'fragrance gift set'
    art = 'An' if fam[:1] in 'aeiou' else 'A'
    leads = [f'Buy {h1} {sz}{pr} at Swansura UK.', f'Buy {h1}{pr} at Swansura UK.', f'{h1}{pr} at Swansura UK.']
    mids = [f' {art} {fam} {gw} {kind}.' if fam else f' A {gw} {kind}.', '']
    tails = [' 100% authentic, fast UK delivery, free over £50.', ' 100% authentic, free UK delivery over £50.', ' Free UK delivery over £50.']
    desc = fit([a + b + c for a in leads for b in mids for c in tails], 140, 160, 'max')
    out.append(dict(p=p, primary=p['Core keyword'], secondary=sec, h1=h1, title=title, desc=desc, note=note))
wb = openpyxl.load_workbook(WB)
if 'Final SEO' in wb.sheetnames: del wb['Final SEO']
ws = wb.create_sheet('Final SEO', 1)
A = 'Arial'; F = Font(name=A, size=10); B = Font(name=A, size=10, bold=True); H = Font(name=A, bold=True, color='FFFFFF', size=10)
HF = PatternFill('solid', fgColor='1F3B34'); BR = PatternFill('solid', fgColor='9A7536')
cols = ['ID', 'Source', 'Brand', 'Product name (sheet)', 'Size', 'Price (£)', 'Primary keyword', 'Primary vol.', 'Primary KD', 'Ease', 'Secondary keyword', 'Secondary vol.', 'FINAL H1', 'FINAL meta title', 'Title len', 'FINAL meta description', 'Desc len', 'Notes']
widths = [9, 15, 16, 44, 8, 8, 32, 9, 8, 9, 36, 9, 46, 60, 7, 90, 7, 50]
for i, (c, w) in enumerate(zip(cols, widths), 1):
    x = ws.cell(row=1, column=i, value=c); x.font = H; x.fill = BR if 'vol.' in c or c in ('Primary KD', 'Ease') else HF; x.alignment = Alignment(wrap_text=True, vertical='center')
    ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
ws.row_dimensions[1].height = 30; ws.freeze_panes = 'E2'
SD = "'Semrush Data'"
look = lambda c, col: f'=IFERROR(INDEX({SD}!${col}$2:${col}$5000,MATCH({c},{SD}!$A$2:$A$5000,0)),"")'
thin = Border(bottom=Side(style='thin', color='E3DACD'))
for r, o in enumerate(out, 2):
    p = o['p']
    vals = [p['ID'], p['Source sheet'], p['Brand'], p['Product name'], p['Size'], p['MRP (£)'], o['primary'], None, None, None, o['secondary'], None, o['h1'], o['title'], None, o['desc'], None, '; '.join(x for x in [o['note'], p['Brand note'] or ''] if x)]
    for c, val in enumerate(vals, 1):
        x = ws.cell(row=r, column=c, value=val); x.font = F; x.border = thin; x.alignment = Alignment(vertical='top', wrap_text=c in (16, 18))
    ws.cell(row=r, column=8, value=look(f'G{r}', 'C')).number_format = '#,##0'
    ws.cell(row=r, column=9, value=look(f'G{r}', 'D'))
    ws.cell(row=r, column=10, value=f'=IF(OR(H{r}="",H{r}=0),"No volume",IF(I{r}<=20,"Easier",IF(I{r}<=35,"Medium","Hard")))')
    ws.cell(row=r, column=12, value=f'=IF(K{r}="","",' + look(f'K{r}', 'C')[1:] + ')').number_format = '#,##0'
    ws.cell(row=r, column=15, value=f'=LEN(N{r})'); ws.cell(row=r, column=17, value=f'=LEN(P{r})')
    ws.cell(row=r, column=6).number_format = '£0.00'
    for c in (8, 9, 10, 12, 15, 17): ws.cell(row=r, column=c).font = F
    for c in (13, 14, 16): ws.cell(row=r, column=c).font = B if c == 13 else F
n = len(out) + 1
red = Font(name=A, size=10, color='C00000', bold=True); grn = Font(name=A, size=10, color='008000', bold=True)
ws.conditional_formatting.add(f'O2:O{n}', FormulaRule(formula=['OR(O2<50,O2>60)'], font=red))
ws.conditional_formatting.add(f'Q2:Q{n}', FormulaRule(formula=['OR(Q2<140,Q2>160)'], font=red))
ws.conditional_formatting.add(f'J2:J{n}', FormulaRule(formula=['J2="Easier"'], font=grn))
ws.conditional_formatting.add(f'R2:R{n}', FormulaRule(formula=['R2<>""'], fill=PatternFill('solid', fgColor='FCE4D6')))
ws.auto_filter.ref = f'A1:R{n}'
wb.save(WB); print('final rows', len(out))
