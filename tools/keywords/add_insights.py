import openpyxl, sys
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.formatting.rule import FormulaRule
WB = sys.argv[1]
vals = openpyxl.load_workbook(WB, data_only=True)
P = [dict(zip([c.value for c in vals['Products'][1]], r)) for r in vals['Products'].iter_rows(min_row=2, values_only=True)]
K = [dict(zip([c.value for c in vals['Keyword List'][1]], r)) for r in vals['Keyword List'].iter_rows(min_row=2, values_only=True)]
num = lambda v: v if isinstance(v, (int, float)) else 0
wb = openpyxl.load_workbook(WB)
if 'Insights' in wb.sheetnames: del wb['Insights']
ws = wb.create_sheet('Insights', 1)
A = 'Arial'; F = Font(name=A, size=10); B = Font(name=A, size=10, bold=True); T = Font(name=A, size=13, bold=True)
H = Font(name=A, bold=True, color='FFFFFF', size=10); HF = PatternFill('solid', fgColor='1F3B34'); G = Font(name=A, size=10, color='008000')
SD = "'Semrush Data'"
def look(cell, col):  # col letter in Semrush Data: C volume, D KD, E CPC, B intent
    return f'=IFERROR(INDEX({SD}!${col}$2:${col}$5000,MATCH({cell},{SD}!$A$2:$A$5000,0)),"")'
for c, w in zip('ABCDEFGH', [10, 52, 40, 11, 9, 9, 16, 46]): ws.column_dimensions[c].width = w
r = 1
def title(t, sub=None):
    global r
    ws.cell(row=r, column=1, value=t).font = T; r += 1
    if sub: c = ws.cell(row=r, column=1, value=sub); c.font = Font(name=A, size=10, italic=True); r += 1
def head(cols):
    global r
    for i, h in enumerate(cols, 1):
        c = ws.cell(row=r, column=i, value=h); c.font = H; c.fill = HF; c.alignment = Alignment(wrap_text=True)
    r += 1
title('Semrush insights — UK, Tiers 1–3', 'Figures are formulas over the "Semrush Data" tab (Semrush Keyword Overview, UK database: Tier 1 batches 1–9, Tier 2 T2-1 to T2-3, Tier 3 T3-1 to T3-2, Oct 2026). Volume = average monthly UK searches; KD = keyword difficulty 0–100; 0 / blank = no Semrush data.')
r += 1
title('Summary')
summary = [
 ('Keywords checked in Semrush', f"=COUNTA({SD}!A2:A5000)"),
 ('Keywords with UK search volume', f'=COUNTIF({SD}!C2:C5000,">0")'),
 ('Products whose primary keyword has volume', '=COUNTIF(Products!' + 'AB2:AB600' + ',">0")'),
 ('…of which Vilvah & Deyga (of 159)', '=COUNTIFS(Products!B2:B600,"Vilvah & Deyga",Products!AB2:AB600,">0")'),
 ('…of which fragrance (of 387)', '=COUNTIFS(Products!B2:B600,"Fragrance Catalogue",Products!AB2:AB600,">0")'),
 ('Total monthly searches for product primary keywords', '=SUM(Products!AB2:AB600)'),
]
# locate Products 'Primary volume' column letter
ph = [c.value for c in wb['Products'][1]]
from openpyxl.utils import get_column_letter as L
pv = L(ph.index('Primary volume') + 1); pk = L(ph.index('Primary keyword') + 1); pkd = L(ph.index('Primary KD') + 1)
for lab, f in summary:
    f = f.replace('AB', pv)
    ws.cell(row=r, column=2, value=lab).font = F; c = ws.cell(row=r, column=4, value=f); c.font = B; c.number_format = '#,##0'; r += 1
r += 1
title('What the data says')
for t in [
 '1. Fragrance carries the search demand: designer lines (Dior Sauvage, Tom Ford, Gucci, Marc Jacobs, Calvin Klein) get thousands of searches a month, but big retailers own page 1, so pick the lower-KD ones (KD ≤ 20) first.',
 '2. Vilvah and Deyga product names get very few searches (10–30 a month each); "vilvah" (480), "vilvah milk drops" (390) and "deyga" (140) are the brand terms people use. Brand pages + the Milk Drops page should target these.',
 '3. The skincare opportunity is non-branded: "moisturiser for dry skin" (2,900), "hair mask for dry hair" (2,400), "hair oil for dry hair" (1,900), "face wash for dry/oily skin" (1,600 each), "shampoo for dandruff" (1,600). Build concern collection pages for these and use the phrases in Vilvah/Deyga product copy and FAQs.',
 '4. Product pages: keep the brand + product name as H1 (it is the primary keyword for 265 products); add the matching concern phrase in the first paragraph and meta description.',
 '5. Tier 2 checked perfume-type / "+ uk" variants for the top 100 fragrances and generic ingredient and product-type terms; the Primary keyword per product now picks the highest-volume variant.']:
    c = ws.cell(row=r, column=2, value=t); c.font = F; c.alignment = Alignment(wrap_text=True, vertical='top'); ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=8); ws.row_dimensions[r].height = 30; r += 1
r += 1
title('Top 50 hero products', 'Ranked by primary-keyword volume, KD as tie-break. Suggested set for the "premium description rewrite × 50" in the quote. Ranking is a snapshot; volumes are live formulas.')
head(['ID', 'Product', 'Primary keyword', 'Volume', 'KD', 'CPC ($)', 'Ease', 'Where to use'])
hero = sorted([p for p in P if num(p['Primary volume']) > 0], key=lambda p: (-num(p['Primary volume']), num(p['Primary KD'])))
seen = set(); top = []
for p in hero:
    if p['Primary keyword'] in seen: continue
    seen.add(p['Primary keyword']); top.append(p)
    if len(top) == 50: break
first = r
for p in top:
    ws.cell(row=r, column=1, value=p['ID']).font = F
    ws.cell(row=r, column=2, value=p['Product name']).font = F
    ws.cell(row=r, column=3, value=f'=INDEX(Products!${pk}$2:${pk}$600,MATCH(A{r},Products!$A$2:$A$600,0))').font = G
    ws.cell(row=r, column=4, value=look(f'C{r}', 'C')).number_format = '#,##0'
    ws.cell(row=r, column=5, value=look(f'C{r}', 'D'))
    ws.cell(row=r, column=6, value=look(f'C{r}', 'E')).number_format = '$0.00'
    ws.cell(row=r, column=7, value=f'=IF(E{r}="","",IF(E{r}<=20,"Easier",IF(E{r}<=35,"Medium","Hard")))')
    ws.cell(row=r, column=8, value='H1 + meta title + first sentence; long-form description (premium rewrite)').font = F
    for c in range(4, 8): ws.cell(row=r, column=c).font = F
    r += 1
ws.conditional_formatting.add(f'G{first}:G{r-1}', FormulaRule(formula=[f'G{first}="Easier"'], font=Font(name=A, size=10, color='008000', bold=True)))
r += 1
def cluster(t, sub, types, use, n=30):
    global r
    title(t, sub); head(['', 'Keyword', 'Type', 'Volume', 'KD', 'CPC ($)', 'Ease', 'Where to use'])
    xs = sorted([k for k in K if k['Type'] in types and num(k['Volume']) > 0], key=lambda k: -num(k['Volume']))[:n]
    for k in xs:
        ws.cell(row=r, column=2, value=k['Keyword']).font = F; ws.cell(row=r, column=3, value=k['Type']).font = F
        ws.cell(row=r, column=4, value=look(f'B{r}', 'C')).number_format = '#,##0'
        ws.cell(row=r, column=5, value=look(f'B{r}', 'D')); ws.cell(row=r, column=6, value=look(f'B{r}', 'E')).number_format = '$0.00'
        ws.cell(row=r, column=7, value=f'=IF(E{r}="","",IF(E{r}<=20,"Easier",IF(E{r}<=35,"Medium","Hard")))')
        ws.cell(row=r, column=8, value=use(k['Keyword'])).font = F
        for c in range(4, 8): ws.cell(row=r, column=c).font = F
        r += 1
    r += 1
cluster('Non-branded skincare & haircare opportunities', 'Category and concern searches: the main organic opportunity for Vilvah & Deyga.', {'Category / concern'},
        lambda k: 'Concern collection page (H1 + intro + FAQ); product copy & meta description' if ' for ' in k else 'Category page H1 / intro; homepage section')
cluster('Ingredient & product-type searches (Tier 2)', 'Generic searches that match the Vilvah & Deyga range: candidates for new collection pages, product-copy phrases and blog topics.', {'Generic product term', 'VD target keyword'},
        lambda k: 'New collection / keyword page + product copy' if not k.endswith(' uk') else 'Homepage / brand page copy', n=45)
cluster('Questions & buying searches (Tier 3)', 'What people ask and how they shop for the hero and Vilvah & Deyga products.', {'Google suggest', 'Buy online', 'Price'},
        lambda k: 'Meta title / description wording' if k.startswith('buy ') or k.endswith(' price') else 'Product FAQ question' if any(w in k for w in ('review', 'how to', 'dupe', 'notes', 'smell', 'benefit', 'vs', 'similar', 'long lasting', 'ingredients')) else 'Product copy / secondary keyword', n=50)
cluster('Vilvah & Deyga brand searches', 'What people type about the two partner brands.', {'Brand', 'Google suggest (brand)'},
        lambda k: 'Brand page title + intro' if k in ('vilvah', 'deyga', 'vilvah uk', 'deyga products', 'vilvah products', 'deyga organics', 'vilvah store', 'vilvah official website') else 'Product page H1 / meta (matching product)')
cluster('Fragrance brand & scent-family searches', 'Brand pages and scent-family collections. High-volume designer terms are hard (LOOKFANTASTIC, Boots, The Perfume Shop); target "Easier" ones first.', {'Brand / family'},
        lambda k: 'Scent-family collection page' if k.split()[0] in ('amber', 'woody', 'floral', 'oud', 'fresh', 'aromatic', 'spicy', 'citrus', 'fruity', 'musk', 'green', 'aquatic', 'leather', 'gourmand', 'chypre') else 'Brand page title + intro', n=40)
ws.freeze_panes = 'A2'
wb.save(WB); print('insights rows', r, 'hero', len(top))
