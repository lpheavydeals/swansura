"""Focused Tier 2: variants for in-demand products + VD target and generic product-type terms. Writes prompts workbook and adds rows to the research workbook."""
import openpyxl, sys, re
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
WB, OUT = sys.argv[1], sys.argv[2]
v = openpyxl.load_workbook(WB, data_only=True)
done = {r[0] for r in v['Semrush Data'].iter_rows(min_row=2, values_only=True) if r[0]}
P = [dict(zip([c.value for c in v['Products'][1]], r)) for r in v['Products'].iter_rows(min_row=2, values_only=True)]
VDC = [dict(zip([c.value for c in v['VD Content'][1]], r)) for r in v['VD Content'].iter_rows(min_row=2, values_only=True)]
num = lambda x: x if isinstance(x, (int, float)) else 0
picked = []  # (keyword, type, product id)
def add(k, t, pid=''):
    k = re.sub(r'\s+', ' ', k.lower().strip())
    if ',' in k or k in done or any(k == x[0] for x in picked): return
    picked.append((k, t, pid))
# A. top 100 in-demand fragrance products: perfume-type + UK variants
fr = sorted([p for p in P if p['Source sheet'] != 'Vilvah & Deyga' and num(p['Primary volume']) > 0], key=lambda p: -num(p['Primary volume']))
seen = set()
for p in fr:
    if p['Core keyword'] in seen: continue
    seen.add(p['Core keyword'])
    if p['+ Concentration']: add(p['+ Concentration'], 'Product + type', p['ID'])
    add(p['+ UK'], 'Product + UK', p['ID'])
    if len(seen) >= 100: break
# B. VD target keywords not yet measured
for r in VDC: add(r['Target keyword'], 'VD target keyword', r['ID'])
# C. generic product-type / ingredient terms that match the Vilvah & Deyga range
GEN = '''goat milk soap|goat milk shampoo|goat milk body lotion|goat milk face wash|rice water shampoo|rice water hair mask|beetroot lip balm|tinted lip balm|lip scrub|natural lip balm|spf lip balm|lip balm for pigmented lips|kumkumadi oil|kumkumadi serum|aloe vera gel|vitamin c serum|vitamin c face wash|niacinamide serum|under eye cream|under eye cream for dark circles|charcoal face wash|charcoal soap|neem face wash|tea tree face wash|onion hair oil|hair growth oil|amla hair oil|castor oil for hair|hibiscus hair oil|bhringraj oil|ayurvedic hair oil|red sandalwood face pack|saffron face cream|multani mitti face pack|ubtan|clay face mask|rose water toner|natural soap bar|handmade soap uk|bath bar|body butter|chocolate body butter|cocoa butter body lotion|foot cream|natural deodorant for women|aluminium free deodorant|hair removal powder|natural hair dye|henna hair dye|indigo powder for hair|ayurvedic skincare uk|indian skincare brands uk|indian skincare uk|milk face wash|oat milk face wash|coffee body scrub|sunscreen spf 50|tinted sunscreen|night cream|face serum for glowing skin|hair mask for hair growth|conditioner for dry hair|shampoo for dry hair|shampoo for hair loss|face wash for sensitive skin|moisturiser for sensitive skin|moisturiser for combination skin|moisturiser for oily skin|toner for oily skin|serum for dark spots|skincare gift set|natural skincare gift set|hair care gift set|vegan skincare uk|cruelty free skincare uk|paraben free shampoo|sulphate free shampoo uk|natural body wash|natural face wash|natural moisturiser|natural hair oil|natural sunscreen uk|organic skincare uk'''.split('|')
for g in GEN: add(g, 'Generic product term')
B = [picked[i:i+100] for i in range(0, len(picked), 100)]
# ---- prompts workbook
A = 'Arial'; H = Font(name=A, bold=True, color='FFFFFF', size=10); F = Font(name=A, size=10); HF = PatternFill('solid', fgColor='1F3B34'); IN = PatternFill('solid', fgColor='FFF2CC')
wb = openpyxl.Workbook(); ix = wb.active; ix.title = 'Index'
ix['A1'] = 'Semrush Tier 2 (focused) prompts for Claude in Chrome'; ix['A1'].font = Font(name=A, bold=True, size=13)
notes = [f'{len(picked)} keywords in {len(B)} batches: perfume-type and "+ uk" variants for the top 100 in-demand fragrances, the Vilvah & Deyga target keywords not yet measured, and generic product/ingredient terms that match the Vilvah & Deyga range.',
         'How to use: open a "T2 Prompt" tab, click the column A header, copy and paste into the Claude in Chrome side panel with Semrush open. One prompt at a time.',
         'Then upload the exports (Tier2_Batch_N_UK_keywords.xlsx) to the Claude Code session.']
for i, n in enumerate(notes, 2): ix.cell(row=i, column=1, value=n).font = F
hdr = ['Prompt', 'Tab', 'Keywords', 'Status', 'Export file name', 'Keywords with data', 'Notes']
for i, h in enumerate(hdr, 1): c = ix.cell(row=6, column=i, value=h); c.font = H; c.fill = HF
for c, w in zip('ABCDEFG', [12, 14, 10, 14, 36, 16, 40]): ix.column_dimensions[c].width = w
dv = DataValidation(type='list', formula1='"To do,Done,Limit hit,Error"', allow_blank=True); ix.add_data_validation(dv)
for bi, batch in enumerate(B, 1):
    lines = f'''Task: get UK search volumes for a keyword list in Semrush (Tier 2, Batch {bi} of {len(B)}).

Steps:
1. In this Semrush tab, open SEO > Keyword Overview.
2. Set the database to United Kingdom (UK flag) before analysing. Do not use any other country.
3. Paste the keyword list below into the Keyword Overview box (bulk analysis, one keyword per line) and click Analyze / Search.
4. When the results table loads, use Export > Excel (.xlsx) to download the results for all keywords. Keep the default columns. Save it as Tier2_Batch_{bi}_UK_keywords.xlsx.
5. Tell me: the downloaded file name, how many keywords returned data, and any keywords Semrush rejected.

Rules:
- Do not change account settings, buy anything, start a trial or upgrade the plan.
- If Semrush shows a limit, captcha, login prompt or an "upgrade" message, stop and tell me instead of working around it.
- Do not add, remove or edit keywords in the list.

Keywords (Tier 2, Batch {bi}):'''.split('\n') + [k for k, _, _ in batch]
    ws = wb.create_sheet(f'T2 Prompt {bi}'); ws.column_dimensions['A'].width = 120
    for r, l in enumerate(lines, 1):
        c = ws.cell(row=r, column=1, value=l or None); c.font = F; c.data_type = 's' if l else c.data_type
    row = 6 + bi
    for i, val in enumerate([f'T2 Prompt {bi}', f'T2 Prompt {bi}', len(batch), 'To do'], 1): ix.cell(row=row, column=i, value=val).font = F
    ix.cell(row=row, column=2).hyperlink = f"#'T2 Prompt {bi}'!A1"; ix.cell(row=row, column=2).font = Font(name=A, size=10, color='0563C1', underline='single')
    dv.add(f'D{row}')
    for c in 'DEFG': ix[f'{c}{row}'].fill = IN
t = 7 + len(B)
ix.cell(row=t, column=2, value='Total').font = Font(name=A, bold=True, size=10)
ix.cell(row=t, column=3, value=f'=SUM(C7:C{t-1})').font = Font(name=A, bold=True, size=10)
ix.cell(row=t, column=4, value=f'=COUNTIF(D7:D{t-1},"Done")&" of {len(B)} done"').font = Font(name=A, bold=True, size=10)
wb.save(OUT)
# ---- add Tier 2 rows to the research workbook (Keyword List + Semrush Batches)
rw = openpyxl.load_workbook(WB)
kl = rw['Keyword List']; existing = {r[0] for r in kl.iter_rows(min_row=2, values_only=True)}
SD = "'Semrush Data'"
look = lambda c, metric: f'=IFERROR(INDEX({SD}!$A$1:$Z$20000,MATCH({c},{SD}!$A$1:$A$20000,0),MATCH("{metric}",{SD}!$A$1:$Z$1,0)),"")'
for bi, batch in enumerate(B, 1):
    for k, typ, pid in batch:
        if k in existing: 
            for r in range(2, kl.max_row + 1):
                if kl.cell(row=r, column=1).value == k: kl.cell(row=r, column=4, value=f'T2-{bi}')
            continue
        r = kl.max_row + 1
        for c, val in enumerate([k, typ, 2, f'T2-{bi}', 1 if pid else 0, pid], 1): kl.cell(row=r, column=c, value=val).font = F
        kl.cell(row=r, column=7, value=look(f'A{r}', 'Volume')).number_format = '#,##0'
        kl.cell(row=r, column=8, value=look(f'A{r}', 'Keyword Difficulty*'))
        kl.cell(row=r, column=9, value=look(f'A{r}', 'CPC*')).number_format = '$0.00'
        kl.cell(row=r, column=10, value=look(f'A{r}', 'Intent'))
        for c in range(7, 11): kl.cell(row=r, column=c).font = F
kl.auto_filter.ref = f'A1:J{kl.max_row}'
sb = rw['Semrush Batches']; base = sb.max_column
for bi, batch in enumerate(B, 1):
    c = sb.cell(row=1, column=base + bi, value=f'T2 Batch {bi}'); c.font = H; c.fill = PatternFill('solid', fgColor='9A7536')
    sb.column_dimensions[openpyxl.utils.get_column_letter(base + bi)].width = 40
    for i, (k, _, _) in enumerate(batch, 2): sb.cell(row=i, column=base + bi, value=k).font = F
rw.save(WB)
from collections import Counter
print('keywords', len(picked), 'batches', len(B), Counter(t for _, t, _ in picked))
