"""Merge Semrush Keyword Overview exports into Swansura_Keyword_Research.xlsx.
Pass ALL exports every time (Tier 1 Batch_N_UK_keywords.xlsx + Tier 2 Tier2_Batch_N_UK_keywords.xlsx): the Semrush Data tab is rebuilt from them."""
import openpyxl, glob, re, sys
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter as L
WB, EXPORTS = sys.argv[1], sys.argv[2:]
A = 'Arial'; F = Font(name=A, size=10); B = Font(name=A, size=10, bold=True); H = Font(name=A, bold=True, color='FFFFFF', size=10); HF = PatternFill('solid', fgColor='1F3B34')
wb = openpyxl.load_workbook(WB)
batches = {}
for col in wb['Semrush Batches'].iter_cols(values_only=True):
    h = str(col[0] or '')
    mt = re.match(r'T(\d) Batch (\d+)', h)
    key = f"T{mt.group(1)}-{mt.group(2)}" if mt else (int(h.split()[1]) if h.startswith('Batch') else None)
    if key is not None: batches[key] = [c for c in col[1:] if c]
rows = {}
for f in EXPORTS:
    m2 = re.search(r'Tier(\d)_Batch_(\d+)', f)
    n = f'T{m2.group(1)}-{m2.group(2)}' if m2 else int(re.search(r'Batch_(\d+)', f).group(1)); want = batches[n]
    data = list(openpyxl.load_workbook(f, data_only=True).worksheets[0].iter_rows(values_only=True))
    hdr = [str(h or '').lower() for h in data[0]]
    col = lambda *names: next((i for i, h in enumerate(hdr) for nm in names if h.startswith(nm)), None)
    ik, ii, iv, ikd, ic, ico, inote = col('keyword'), col('intent'), col('volume'), col('kd', 'keyword difficulty'), col('cpc'), col('competition', 'com'), col('note')
    for r in data[1:]:
        if not r[ik]: continue
        kw = str(r[ik]).strip().lower(); note = (r[inote] if inote is not None and inote < len(r) else None) or ''
        num = lambda v: v if isinstance(v, (int, float)) else None
        rec = dict(kw=kw, intent=r[ii], vol=num(r[iv]) if num(r[iv]) is not None else 0, kd=num(r[ikd]), cpc=num(r[ic]), comp=num(r[ico]), batch=n, note='No Semrush data (n/a)' if num(r[iv]) is None else '')
        if kw not in want:  # comma-split keyword: map back to the original list keyword
            orig = next((w for w in want if w.startswith(kw) and ',' in w), None)
            if orig:
                rows[orig] = dict(rec, kw=orig, note=f'Semrush ran "{kw}" (it splits at commas); figures are for that shorter phrase')
            continue
        if kw in rows and rows[kw]['vol'] >= rec['vol']: continue
        rows[kw] = rec
ws = wb['Semrush Data']; ws.delete_rows(1, ws.max_row)
for c in range(1, 12): ws.cell(row=1, column=c).value = None
hdr = ['Keyword','Intent','Volume','Keyword Difficulty','CPC (USD)','Competition','Batch','Note']
for i, h in enumerate(hdr, 1):
    c = ws.cell(row=1, column=i, value=h); c.font = H; c.fill = HF
for r, rec in enumerate(sorted(rows.values(), key=lambda x: (-x['vol'], x['kw'])), 2):
    for c, v in enumerate([rec['kw'], rec['intent'], rec['vol'], rec['kd'], rec['cpc'], rec['comp'], rec['batch'], rec['note']], 1):
        x = ws.cell(row=r, column=c, value=v); x.font = F; x.fill = PatternFill(fill_type=None)
    ws.cell(row=r, column=3).number_format = '#,##0'; ws.cell(row=r, column=5).number_format = '$0.00'
for c, w in zip('ABCDEFGH', [52,24,10,10,10,11,7,70]): ws.column_dimensions[c].width = w
ws.freeze_panes = 'A2'; ws.auto_filter.ref = f'A1:H{len(rows)+1}'
# currency is USD in this Semrush account
p = wb['Products']; ph = [c.value for c in p[1]]
pc = next(i for i, h in enumerate(ph) if str(h).startswith('Primary CPC')) + 1
for r in range(2, p.max_row + 1): p.cell(row=r, column=pc).number_format = '$0.00'
p.cell(row=1, column=pc).value = 'Primary CPC ($)'
k = wb['Keyword List']; k.cell(row=1, column=9).value = 'CPC ($)'
for r in range(2, k.max_row + 1): k.cell(row=r, column=9).number_format = '$0.00'
wb.save(WB); print('merged', len(rows), 'with volume', sum(1 for x in rows.values() if x['vol'] > 0))
