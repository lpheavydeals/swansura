import json, sys, re, unicodedata
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
KW, AC, OUT = sys.argv[1], sys.argv[2], sys.argv[3]
d = json.load(open(KW)); ac = json.load(open(AC))
P, K = d['products'], {k['kw']:k for k in d['keywords']}
def plain(s): return ''.join(c for c in unicodedata.normalize('NFD', s) if unicodedata.category(c)!='Mn')
# --- relevant autocomplete suggestions -> tier 3 'Google suggest'
STOP = {'perfume','for','the','de','and','&','uk','men','women','eau','parfum','by','of'}
def toks(s): return [t for t in re.findall(r"[a-z0-9']+", plain(s.lower())) if t not in STOP]
MODS = set('review reviews price uk online buy sale discount original genuine perfume parfum extrait edp edt tester gift set sets notes smell smells like dupe how to use benefits ingredients results before after kit combo cream oil serum lotion soap wash balm ml mini travel size sizes amazon india boots'.split())
sugg_by_pid = {}
for p in P:
    core = p['core']; ct = toks(core); bt = toks(p['brand'])[:1]
    out = []
    for s in ac.get(core, []):
        s = plain(s.lower()).strip()
        if s == core or len(s) > 80: continue
        st = set(toks(s))
        extra = st - set(ct) - set(toks(p['brand']))
        if len(set(ct) & st) >= min(2, len(ct)) and (not bt or bt[0] in st or p['brand'].lower() in ('generic','swansura - the body studio')) and all(t in MODS or re.fullmatch(r'\d+(ml|g)?', t) for t in extra):
            out.append(s)
    sugg_by_pid[p['id']] = out[:6]
    for s in out[:6]:
        e = K.setdefault(s, {'kw':s,'kind':'Google suggest','tier':3,'pids':[]})
        if p['id'] not in e['pids']: e['pids'].append(p['id'])
for seed in ['vilvah','deyga','vilvah uk','deyga uk','buy vilvah','buy deyga']:
    for s in ac.get(seed, []):
        s = plain(s.lower()).strip()
        if s.split()[0].rstrip("'s") in ('vilvah','deyga','buy') and s not in K: K[s] = {'kw':s,'kind':'Google suggest (brand)','tier':1,'pids':[]}
kwlist = sorted(K.values(), key=lambda k:(k['tier'], k['kind'], k['kw']))
for i, k in enumerate(kwlist): k['batch'] = i // 100 + 1
NB = kwlist[-1]['batch']

A = 'Arial'; H = Font(name=A, bold=True, color='FFFFFF', size=10); F = Font(name=A, size=10); B = Font(name=A, size=10, bold=True)
HF = PatternFill('solid', fgColor='1F3B34'); IN = PatternFill('solid', fgColor='FFF2CC'); SUB = PatternFill('solid', fgColor='F2EDE4')
thin = Side(style='thin', color='E3DACD'); BD = Border(bottom=thin)
wb = Workbook()
def header(ws, cols, widths):
    for i, (c, w) in enumerate(zip(cols, widths), 1):
        x = ws.cell(row=1, column=i, value=c); x.font = H; x.fill = HF; x.alignment = Alignment(vertical='center', wrap_text=True)
        ws.column_dimensions[L(i)].width = w
    ws.row_dimensions[1].height = 32; ws.freeze_panes = 'A2'
SD = "'Semrush Data'"
def look(cell, metric):  # metric header text (wildcards allowed)
    return f'=IFERROR(INDEX({SD}!$A$1:$Z$20000,MATCH({cell},{SD}!$A$1:$A$20000,0),MATCH("{metric}",{SD}!$A$1:$Z$1,0)),"")'

# ---------- How to use ----------
ws = wb.active; ws.title = 'How to use'
rows = [
 ('Swansura keyword research — product pages', B),
 (f'Built 18 Oct 2026 from "Swansura_Vilvah_Deyga_Fragrance_Data_Sheet" ({len(P)} products: {sum(p["sheet"]=="Vilvah & Deyga" for p in P)} Vilvah & Deyga, {sum(p["sheet"]!="Vilvah & Deyga" for p in P)} fragrance). Keywords are generated from each product name plus real Google UK autocomplete suggestions (google.co.uk, en-GB, collected 18 Oct 2026).', F),
 ('', F),
 ('How to add search volumes (Semrush)', B),
 ('1. Open the "Semrush Batches" tab. Each column is one batch of up to 100 keywords, in priority order (Tier 1 first).', F),
 ('2. In Semrush go to Keyword Overview, set the database to United Kingdom, paste one batch and click Analyze.', F),
 ('3. Export the results (Excel/CSV), then paste the export INCLUDING its header row into the "Semrush Data" tab, starting at cell A1. Add later batches below the previous ones (skip their header rows).', F),
 ('4. Volume, Keyword Difficulty, CPC and Intent fill in automatically on "Products" and "Keyword List". "Primary keyword" picks the highest-volume variant for each product.', F),
 ('Tip: Tier 1 alone (batches 1–' + str(max(k['batch'] for k in kwlist if k['tier']==1)) + ') covers every product\'s core keyword plus brand, category and concern terms — start there.', F),
 ('', F),
 ('Tiers', B),
 ('Tier 1 — core product keyword (brand + product), brand terms, category and concern terms (e.g. "face wash for dry skin").', F),
 ('Tier 2 — product + "uk", product + concentration (e.g. "... eau de parfum").', F),
 ('Tier 3 — commercial modifiers ("buy ... online", "... price") and Google autocomplete suggestions.', F),
 ('', F),
 ('Where to use the keywords', B),
 ('Primary keyword → H1 and the start of the meta title. Product + UK / buy online → meta description and first sentence of the product copy. Category / concern terms → collection pages, product FAQs and internal links. Autocomplete phrases → FAQ questions and "how to use" copy.', F),
 ('The draft H1, meta title and meta description columns on "Products" follow the 50–60 / 140–160 character targets in your sheet; length columns turn red when outside the range.', F),
 ('', F),
 ('Notes and assumptions', B),
 ('• Keywords are lower-case and accent-free, as people type them. Sizes (100ml), "for women", "luxury" etc. are removed from the core keyword.', F),
 ('• "Brand note" flags catalogue issues found while building keywords: products filed under Generic / Swansura - The Body Studio, brand spellings (e.g. Tomford → Tom Ford) and names that mention a different brand than the Brand column.', F),
 ('• Volumes are blank until Semrush data is pasted in. Cells show "" (blank) for keywords Semrush has no data for — treat those as zero/very low volume.', F),
 ('• Meta descriptions use "free UK delivery over £50" and "100% authentic" from the current site; confirm before publishing.', F),
]
for i, (t, f) in enumerate(rows, 1):
    c = ws.cell(row=i, column=1, value=t); c.font = Font(name=A, size=13 if i == 1 else 10, bold=f.bold); c.alignment = Alignment(wrap_text=True, vertical='top')
ws.column_dimensions['A'].width = 130

# ---------- Products ----------
ws = wb.create_sheet('Products')
VAR = ['Core keyword','+ UK','Buy online','Price','+ Concentration']
cols = ['ID','Source sheet','S.No','Brand','Brand note','Category','Type','Concern / scent','Product name','Size','MRP (£)'] + VAR + [f'Vol: {v}' for v in VAR] + ['Primary keyword','Primary volume','Primary KD','Primary CPC','Category / brand keywords','Google autocomplete phrases','Draft H1','Draft meta title','Title length','Draft meta description','Description length']
widths = [9,15,6,18,24,11,18,22,44,8,8] + [34,34,38,34,40] + [9]*5 + [36,9,8,8,46,60,46,60,8,80,8]
header(ws, cols, widths)
ci = {c:i+1 for i,c in enumerate(cols)}
def disp(p):
    n = re.split(r'\s[–—|]\s|\s[–—]|–', p['name'])[0]
    if p['sheet'] != 'Vilvah & Deyga': n = re.sub(r'\s-\s.*$', '', n)
    n = re.sub(r'\b\d+(\.\d+)?\s*(ml|g|gm|ML|G|GM)\b', '', n, flags=re.I); n = re.sub(r'\+.*$', '', n); n = re.sub(r'\s+', ' ', n).strip(' -,')
    if p['sheet'] == 'Vilvah & Deyga' and not n.lower().startswith(p['brand'].lower()): n = f"{p['brand']} {n[:1].upper()+n[1:]}"
    return n
SMALL = {'and','or','of','de','di','du','la','le','for','the','with','in','a','&','pour','des','da'}
def tcase(n):
    w = n.split(' ')
    return ' '.join(x if (x[:1].isupper() or any(ch.isdigit() for ch in x) or (i and x.lower() in SMALL)) else x[:1].upper()+x[1:] for i, x in enumerate(w))
TYPEW = {'Bath & Body Wash':'body wash','Body Butters':'body butter','Body Cream & Lotions':'body lotion','Body Oil':'body oil','Body Scrub & Exfoliators':'body scrub','Cleansers':'face wash','Conditioner':'conditioner','Deodorant':'deodorant','Face Masks':'face mask','Gift Sets & Combos':'gift set','Hair Masks':'hair mask','Hair Oil':'hair oil','Hair Serums':'hair serum','Lip Care':'lip care','Moisturisers':'moisturiser','SPF for Face':'sunscreen','Serums':'serum','Shampoo':'shampoo','Toners':'toner','Skincare Treatments':'skincare treatment','Body Care':'body care','Hair Treatments':'hair treatment'}
CONW = {'Anti-Ageing':'ageing skin','Anti-Dandruff':'a flake-free scalp','Dark Circles':'tired under-eyes','Dry Lips':'dry lips','Dry Skin':'dry skin','Dull Skin':'dull skin','Frizzy Hair':'frizzy hair','Hair Thinning & Loss':'thinning hair','Oily Skin':'oily skin','Acne & Blemishes':'blemish-prone skin','Body Hydration':'soft, hydrated skin','Skin Nourishment':'nourished skin','Hair Nourishment':'nourished hair','Sun Protection':'daily sun protection','Damaged Skin Barrier':'a stressed skin barrier','Dull & Uneven Skin':'uneven skin tone'}
for r, p in enumerate(P, 2):
    v = p['kws'] + [''] * (5 - len(p['kws']))
    if p['sheet'] != 'Vilvah & Deyga':  # fragrance list = core, conc?, uk, buy, price
        conc = [k for k in p['kws'][1:] if not k.endswith(' uk') and not k.startswith('buy ') and not k.endswith(' price')]
        v = [p['core'], f"{p['core']} uk", f"buy {p['core']} online", f"{p['core']} price", conc[0] if conc else '']
    h1 = tcase(disp(p))
    size = (str(p['size']).replace(' ','') if p['size'] and str(p['size']).lower()!='none' else '')
    tw = TYPEW.get(p['type'], (p['type'] or '').lower())
    if p['sheet'] == 'Vilvah & Deyga':
        tcands = [f"{h1} | Buy Online UK | Swansura", f"{h1} {size} | Buy Online UK | Swansura" if size else '', f"{h1} | Official UK Partner | Swansura", f"Buy {h1} {size} Online UK | Swansura" if size else '', f"{h1} | Swansura UK"]
    else:
        conc = {'Eau de Parfum (EDP)':'EDP','Eau de Toilette (EDT)':'EDT','Extrait de Parfum':'Extrait','Parfum':'Parfum'}.get(p['type'],'')
        h1b = re.sub(r'\s*\b(Extrait de Parfum|Eau de Parfum|Eau de Toilette|Eau de Cologne|EDP|EDT|Parfum Intense|Parfum)\b\s*', ' ', h1, flags=re.I).strip()
        tcands = [f"{h1} | Buy Online UK | Swansura", f"{h1} {size} | Buy Online UK | Swansura" if size else '', f"{h1b} {conc} {size} | Buy Online UK | Swansura" if size and conc else '', f"{h1} | Swansura UK", f"{h1b} {conc} {size} | Swansura UK" if conc else '', f"{h1b} | Buy Online UK | Swansura", f"{h1b} {size} | Swansura UK" if size else '', f"{h1b} | Swansura UK"]
    tcands = [re.sub(r'\s+',' ',t) for t in tcands if t]
    def cut(t):
        w = t[:46].rsplit(' ',1)[0].split(' ')
        while w and w[-1].lower() in SMALL | {'-','–'}: w.pop()
        return ' '.join(w) + ' | Swansura UK'
    title = next((t for t in tcands if 50 <= len(t) <= 60), None) or max((t for t in tcands if len(t) <= 60), key=len, default=None) or cut(h1b if p['sheet'] != 'Vilvah & Deyga' else h1)
    if p['sheet'] == 'Vilvah & Deyga':
        con = CONW.get(p['concern'])
        leads = [f"Buy {h1} online in the UK at Swansura, official UK partner of {p['brand']}.", f"Buy {h1} online at Swansura, official UK partner of {p['brand']}."]
        mids = [f" A natural {tw} for {con}." if con and tw else '', f" Natural {tw}." if tw else '', '']
        tails = [" 100% authentic, shipped from the UK, free delivery over £50.", " 100% authentic, free UK delivery over £50.", " Free UK delivery over £50.", '']
    else:
        g = {'Unisex':'unisex','Women':"women's",'Men':"men's"}
        gg = next((g[k] for k in g if k in (p['concern'] or '')), '')
        fam = (p['concern'] or '').split(' · ')[0].lower()
        cw = {'Eau de Parfum (EDP)':'eau de parfum','Eau de Toilette (EDT)':'eau de toilette','Extrait de Parfum':'extrait de parfum','Parfum':'parfum'}.get(p['type'],'fragrance')
        leads = [f"Buy {h1} online in the UK at Swansura.", f"Buy {h1} at Swansura UK."]
        mids = [f" A {fam} {gg} {cw}{', '+size if size else ''}.", f" A {fam} {gg} {cw}.", f" {fam.title()} {cw}.", '']
        tails = [" 100% authentic, fast UK delivery, free over £50.", " 100% authentic, free UK delivery over £50.", " Free UK delivery over £50.", '']
    combos = [re.sub(r'\bA ([aeiouAEIOU])', r'An \1', re.sub(r'\s+',' ', (a+b+c))).strip() for a in leads for b in mids for c in tails]
    ok = [x for x in combos if 140 <= len(x) <= 160]
    desc = (max(ok, key=len) if ok else min(combos, key=lambda x: abs(len(x)-150) if len(x) <= 160 else 999))
    vals = [p['id'], p['sheet'], p['row'], p['brand'], p['brand_fix'], p['cat'], p['type'], p['concern'], p['name'], p['size'], p['price']] + v
    for c, val in enumerate(vals, 1): ws.cell(row=r, column=c, value=val)
    k0 = ci['Core keyword']; vcol = ci['Vol: Core keyword']
    for j in range(5):
        kc = f'{L(k0+j)}{r}'
        ws.cell(row=r, column=vcol+j, value=f'=IF({kc}="","",' + look(kc, 'Volume')[1:] + ')')
    kr, vr = f'{L(k0)}{r}:{L(k0+4)}{r}', f'{L(vcol)}{r}:{L(vcol+4)}{r}'
    ws.cell(row=r, column=ci['Primary keyword'], value=f'=IF(COUNT({vr})=0,{L(k0)}{r},IF(MAX({vr})=0,{L(k0)}{r},INDEX({kr},MATCH(MAX({vr}),{vr},0))))')
    pk = f"{L(ci['Primary keyword'])}{r}"
    ws.cell(row=r, column=ci['Primary volume'], value=look(pk, 'Volume'))
    ws.cell(row=r, column=ci['Primary KD'], value=look(pk, 'Keyword Difficulty*'))
    ws.cell(row=r, column=ci['Primary CPC'], value=look(pk, 'CPC*'))
    ws.cell(row=r, column=ci['Category / brand keywords'], value=' · '.join(p['cluster']))
    ws.cell(row=r, column=ci['Google autocomplete phrases'], value=' · '.join(sugg_by_pid.get(p['id'], [])))
    ws.cell(row=r, column=ci['Draft H1'], value=h1)
    ws.cell(row=r, column=ci['Draft meta title'], value=title)
    ws.cell(row=r, column=ci['Title length'], value=f"=LEN({L(ci['Draft meta title'])}{r})")
    ws.cell(row=r, column=ci['Draft meta description'], value=desc)
    ws.cell(row=r, column=ci['Description length'], value=f"=LEN({L(ci['Draft meta description'])}{r})")
    for c in range(1, len(cols)+1):
        x = ws.cell(row=r, column=c); x.font = F; x.border = BD
        if cols[c-1].startswith('Vol') or cols[c-1].startswith('Primary ') and cols[c-1] != 'Primary keyword': x.number_format = '#,##0'
    ws.cell(row=r, column=ci['Primary CPC']).number_format = '£0.00'
    ws.cell(row=r, column=ci['Primary keyword']).font = B
n = len(P) + 1
red = Font(name=A, size=10, color='C00000', bold=True)
tl, dl = L(ci['Title length']), L(ci['Description length'])
ws.conditional_formatting.add(f'{tl}2:{tl}{n}', FormulaRule(formula=[f'OR({tl}2<50,{tl}2>60)'], font=red))
ws.conditional_formatting.add(f'{dl}2:{dl}{n}', FormulaRule(formula=[f'OR({dl}2<140,{dl}2>160)'], font=red))
bn = L(ci['Brand note']); ws.conditional_formatting.add(f'{bn}2:{bn}{n}', FormulaRule(formula=[f'{bn}2<>""'], fill=PatternFill('solid', fgColor='FCE4D6')))
for c in range(ci['Vol: Core keyword'], ci['Primary CPC']+1): ws.cell(row=1, column=c).fill = PatternFill('solid', fgColor='9A7536')
ws.auto_filter.ref = f'A1:{L(len(cols))}{n}'

# ---------- Keyword List ----------
ws = wb.create_sheet('Keyword List')
cols = ['Keyword','Type','Tier','Semrush batch','No. of products','Product IDs','Volume','Keyword Difficulty','CPC (£)','Intent']
header(ws, cols, [52,24,6,9,9,40,10,10,9,16])
for r, k in enumerate(kwlist, 2):
    for c, val in enumerate([k['kw'], k['kind'], k['tier'], k['batch'], len(k['pids']), ', '.join(k['pids'][:12]) + (' …' if len(k['pids'])>12 else '')], 1): ws.cell(row=r, column=c, value=val).font = F
    ws.cell(row=r, column=7, value=look(f'A{r}', 'Volume')).number_format = '#,##0'
    ws.cell(row=r, column=8, value=look(f'A{r}', 'Keyword Difficulty*'))
    ws.cell(row=r, column=9, value=look(f'A{r}', 'CPC*')).number_format = '£0.00'
    ws.cell(row=r, column=10, value=look(f'A{r}', 'Intent'))
    for c in range(7, 11): ws.cell(row=r, column=c).font = F
for c in range(7, 11): ws.cell(row=1, column=c).fill = PatternFill('solid', fgColor='9A7536')
ws.auto_filter.ref = f'A1:J{len(kwlist)+1}'

# ---------- Semrush Batches ----------
ws = wb.create_sheet('Semrush Batches')
for b in range(1, NB+1):
    items = [k['kw'] for k in kwlist if k['batch'] == b]; tiers = sorted({k['tier'] for k in kwlist if k['batch'] == b})
    c = ws.cell(row=1, column=b, value=f'Batch {b} · Tier {"/".join(map(str,tiers))}'); c.font = H; c.fill = HF
    ws.column_dimensions[L(b)].width = 40
    for i, kw in enumerate(items, 2): ws.cell(row=i, column=b, value=kw).font = F
ws.freeze_panes = 'A2'; ws.row_dimensions[1].height = 24

# ---------- Semrush Data (paste area) ----------
ws = wb.create_sheet('Semrush Data')
cols = ['Keyword','Intent','Volume','Keyword Difficulty','CPC (GBP)','Competitive Density','SERP Features','Number of Results']
header(ws, cols, [52,16,10,10,10,10,30,14])
note = ws.cell(row=1, column=10, value='Paste the Semrush Keyword Overview export here, header row included, from A1 (replace this header). Later batches: paste below, without their header. Column order does not matter as long as Keyword is in column A.')
note.font = Font(name=A, size=10, italic=True); note.alignment = Alignment(wrap_text=True, vertical='top'); ws.column_dimensions['J'].width = 60
for r in range(2, 60):
    for c in range(1, 9): ws.cell(row=r, column=c).fill = IN

# ---------- Autocomplete ----------
ws = wb.create_sheet('Google Autocomplete')
header(ws, ['Seed query (google.co.uk)','Suggestion','Used in keyword list'], [48,60,12])
r = 2
for seed in sorted(ac):
    for s in ac[seed]:
        s2 = plain(s.lower()).strip()
        for c, val in enumerate([seed, s2, 'Yes' if s2 in K else ''], 1): ws.cell(row=r, column=c, value=val).font = F
        r += 1
ws.auto_filter.ref = f'A1:C{r-1}'
wb.save(OUT)
print('keywords', len(kwlist), 'batches', NB, 'tier1 batches', max(k['batch'] for k in kwlist if k['tier']==1), 'autocomplete rows', r-2)
