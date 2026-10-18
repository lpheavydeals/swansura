"""Small Tier 3: question-style autocomplete phrases (top 50 hero + Vilvah/Deyga) and buy/price for the top 20 heroes."""
import openpyxl, sys, re
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
WB, OUT = sys.argv[1], sys.argv[2]
v = openpyxl.load_workbook(WB, data_only=True)
done = {r[0] for r in v['Semrush Data'].iter_rows(min_row=2, values_only=True) if r[0]}
P = [dict(zip([c.value for c in v['Products'][1]], r)) for r in v['Products'].iter_rows(min_row=2, values_only=True)]
num = lambda x: x if isinstance(x, (int, float)) else 0
ranked = sorted([p for p in P if num(p['Primary volume']) > 0], key=lambda p: (-num(p['Primary volume']), num(p['Primary KD'])))
seen, top = set(), []
for p in ranked:
    if p['Primary keyword'] in seen: continue
    seen.add(p['Primary keyword']); top.append(p)
    if len(top) == 50: break
PRI = ['review', 'dupe', 'how to use', 'benefits', 'notes', 'smell', 'similar', 'vs', 'long lasting', 'ingredients', 'results', 'original', 'tester', 'sample', 'gift set']
def score(x): return next((i for i, w in enumerate(PRI) if w in x), 99)
def ac(p): return sorted([x for x in (p['Google autocomplete phrases'] or '').split(' · ') if x and ',' not in x and x not in done and not x.endswith(' price')], key=score)
picked = []
def add(k, t, pid):
    k = re.sub(r'\s+', ' ', k.lower().strip())
    if k and ',' not in k and k not in done and all(k != x[0] for x in picked): picked.append((k, t, pid))
for p in top[:20]:
    add(f"buy {p['Core keyword']} online", 'Buy online', p['ID']); add(f"{p['Core keyword']} price", 'Price', p['ID'])
for p in top:
    for x in ac(p)[:2]: add(x, 'Google suggest', p['ID'])
for p in [p for p in P if p['Source sheet'] == 'Vilvah & Deyga']:
    for x in ac(p)[:1]: add(x, 'Google suggest', p['ID'])
B = [picked[i:i+100] for i in range(0, len(picked), 100)]
A = 'Arial'; H = Font(name=A, bold=True, color='FFFFFF', size=10); F = Font(name=A, size=10); HF = PatternFill('solid', fgColor='1F3B34'); IN = PatternFill('solid', fgColor='FFF2CC')
wb = openpyxl.Workbook(); ix = wb.active; ix.title = 'Index'
ix['A1'] = 'Semrush Tier 3 (small) prompts for Claude in Chrome'; ix['A1'].font = Font(name=A, bold=True, size=13)
notes = [f'{len(picked)} keywords in {len(B)} batches: question-style Google autocomplete phrases (review, dupe, how to use, notes…) for the top 50 hero products and Vilvah & Deyga products, plus "buy … online" and "… price" for the top 20 hero products. Used for product FAQs and meta wording.',
         'How to use: open a "T3 Prompt" tab, click the column A header, copy and paste into the Claude in Chrome side panel with Semrush open. One prompt at a time.',
         'Then upload the exports (Tier3_Batch_N_UK_keywords.xlsx) to the Claude Code session.']
for i, n in enumerate(notes, 2): ix.cell(row=i, column=1, value=n).font = F
for i, h in enumerate(['Prompt', 'Tab', 'Keywords', 'Status', 'Export file name', 'Keywords with data', 'Notes'], 1):
    c = ix.cell(row=6, column=i, value=h); c.font = H; c.fill = HF
for c, w in zip('ABCDEFG', [12, 14, 10, 14, 36, 16, 40]): ix.column_dimensions[c].width = w
dv = DataValidation(type='list', formula1='"To do,Done,Limit hit,Error"', allow_blank=True); ix.add_data_validation(dv)
for bi, batch in enumerate(B, 1):
    lines = f'''Task: get UK search volumes for a keyword list in Semrush (Tier 3, Batch {bi} of {len(B)}).

Steps:
1. In this Semrush tab, open SEO > Keyword Overview.
2. Set the database to United Kingdom (UK flag) before analysing. Do not use any other country.
3. Paste the keyword list below into the Keyword Overview box (bulk analysis, one keyword per line) and click Analyze / Search.
4. When the results table loads, use Export > Excel (.xlsx) to download the results for all keywords. Keep the default columns. Save it as Tier3_Batch_{bi}_UK_keywords.xlsx.
5. Tell me: the downloaded file name, how many keywords returned data, and any keywords Semrush rejected.

Rules:
- Do not change account settings, buy anything, start a trial or upgrade the plan.
- If Semrush shows a limit, captcha, login prompt or an "upgrade" message, stop and tell me instead of working around it.
- Do not add, remove or edit keywords in the list.

Keywords (Tier 3, Batch {bi}):'''.split('\n') + [k for k, _, _ in batch]
    ws = wb.create_sheet(f'T3 Prompt {bi}'); ws.column_dimensions['A'].width = 120
    for r, l in enumerate(lines, 1):
        c = ws.cell(row=r, column=1, value=l or None); c.font = F
        if l: c.data_type = 's'
    row = 6 + bi
    for i, val in enumerate([f'T3 Prompt {bi}', f'T3 Prompt {bi}', len(batch), 'To do'], 1): ix.cell(row=row, column=i, value=val).font = F
    ix.cell(row=row, column=2).hyperlink = f"#'T3 Prompt {bi}'!A1"; ix.cell(row=row, column=2).font = Font(name=A, size=10, color='0563C1', underline='single')
    dv.add(f'D{row}')
    for c in 'DEFG': ix[f'{c}{row}'].fill = IN
t = 7 + len(B)
ix.cell(row=t, column=2, value='Total').font = Font(name=A, bold=True, size=10)
ix.cell(row=t, column=3, value=f'=SUM(C7:C{t-1})').font = Font(name=A, bold=True, size=10)
ix.cell(row=t, column=4, value=f'=COUNTIF(D7:D{t-1},"Done")&" of {len(B)} done"').font = Font(name=A, bold=True, size=10)
wb.save(OUT)
rw = openpyxl.load_workbook(WB)
kl = rw['Keyword List']; rowof = {kl.cell(row=r, column=1).value: r for r in range(2, kl.max_row + 1)}
SD = "'Semrush Data'"
look = lambda c, metric: f'=IFERROR(INDEX({SD}!$A$1:$Z$20000,MATCH({c},{SD}!$A$1:$A$20000,0),MATCH("{metric}",{SD}!$A$1:$Z$1,0)),"")'
for bi, batch in enumerate(B, 1):
    for k, typ, pid in batch:
        if k in rowof: kl.cell(row=rowof[k], column=4, value=f'T3-{bi}'); continue
        r = kl.max_row + 1
        for c, val in enumerate([k, typ, 3, f'T3-{bi}', 1, pid], 1): kl.cell(row=r, column=c, value=val).font = F
        kl.cell(row=r, column=7, value=look(f'A{r}', 'Volume')).number_format = '#,##0'
        kl.cell(row=r, column=8, value=look(f'A{r}', 'Keyword Difficulty*'))
        kl.cell(row=r, column=9, value=look(f'A{r}', 'CPC*')).number_format = '$0.00'
        kl.cell(row=r, column=10, value=look(f'A{r}', 'Intent'))
        for c in range(7, 11): kl.cell(row=r, column=c).font = F
kl.auto_filter.ref = f'A1:J{kl.max_row}'
sb = rw['Semrush Batches']; base = sb.max_column
for bi, batch in enumerate(B, 1):
    c = sb.cell(row=1, column=base + bi, value=f'T3 Batch {bi}'); c.font = H; c.fill = PatternFill('solid', fgColor='3A2421')
    sb.column_dimensions[openpyxl.utils.get_column_letter(base + bi)].width = 40
    for i, (k, _, _) in enumerate(batch, 2): sb.cell(row=i, column=base + bi, value=k).font = F
rw.save(WB)
from collections import Counter
print('keywords', len(picked), 'batches', len(B), Counter(t for _, t, _ in picked))
