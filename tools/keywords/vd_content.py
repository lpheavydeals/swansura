"""Add a 'VD Content' tab: concern-keyword-led H1/meta/intro copy for Vilvah & Deyga products."""
import openpyxl, json, re, sys
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import FormulaRule
WB, PJ = sys.argv[1], sys.argv[2]
vals = openpyxl.load_workbook(WB, data_only=True)
P = [dict(zip([c.value for c in vals['Products'][1]], r)) for r in vals['Products'].iter_rows(min_row=2, values_only=True)]
VD = [p for p in P if p['Source sheet'] == 'Vilvah & Deyga']
pj = [x for x in json.load(open(PJ)) if x['brand'] in ('Vilvah', 'Deyga')]
def toks(s): return set(re.sub(r'[^a-z0-9 ]', '', re.sub(r'\b\d+\s*(ml|g|gm)\b', '', s.lower())).replace('deygas', 'deyga').split())
def match(p):
    t = toks(p['Product name']) | {p['Brand'].lower()}
    best = max(pj, key=lambda x: len(t & toks(x['title'])) / max(1, len(t | toks(x['title']))))
    sc = len(t & toks(best['title'])) / max(1, len(t | toks(best['title'])))
    return best if sc >= 0.6 and best['brand'] == p['Brand'] else None
# (Product Type, Concern) -> (target keyword, prototype page slug)
def target(typ, con):
    con = con or ''
    if typ == 'Moisturisers': return ('moisturiser for dry skin', 'moisturiser-for-dry-skin')
    if typ == 'Cleansers': return ('face wash for oily skin', 'face-wash-for-oily-skin') if con in ('Oily Skin', 'Acne & Blemishes', 'Blackheads & Pores') else ('face wash for dry skin', 'face-wash-for-dry-skin')
    if typ == 'Toners': return ('toner for dry skin', 'toner-for-dry-skin')
    if typ == 'Lip Care': return ('lip balm for dry lips', 'lip-balm-for-dry-lips')
    if typ == 'Body Butters': return ('body butter for dry skin', 'body-butter-for-dry-skin')
    if typ in ('Body Cream & Lotions', 'Body Care', 'Body Oil'): return ('body lotion for dry skin', 'body-lotion-for-dry-skin')
    if typ == 'Hair Oil': return ('hair oil for hair loss', 'hair-oil-for-hair-loss') if 'Loss' in con else ('hair oil for dry hair', 'hair-oil-for-dry-hair')
    if typ == 'Shampoo': return ('shampoo for dandruff', 'shampoo-for-dandruff') if 'Dandruff' in con else ('natural shampoo', 'natural-shampoo')
    if typ in ('Hair Masks', 'Hair Treatments', 'Conditioner'): return ('hair mask for dry hair', 'hair-mask-for-dry-hair')
    if typ == 'Hair Serums': return ('hair serum for frizzy hair', 'hair-serum-for-frizzy-hair')
    if typ == 'SPF for Face': return ('natural sunscreen', 'natural-sunscreen')
    if typ == 'Face Masks': return ('face mask for oily skin', 'face-mask-for-oily-skin') if con == 'Oily Skin' else ('face mask', 'face-mask')
    if typ in ('Serums', 'Skincare Treatments'): return ('skincare for dull skin', 'dull-skin') if 'Dull' in con or 'Pigment' in con else ('face serum', 'face-serum')
    if typ == 'Bath & Body Wash': return ('natural body wash', 'shower-gel-body-wash')
    if typ == 'Deodorant': return ('natural deodorant', 'deodorant')
    if typ == 'Body Scrub & Exfoliators': return ('body scrub', 'body-scrub')
    if typ == 'Gift Sets & Combos':
        m = {'Dry Skin': ('skincare for dry skin', 'dry-skin'), 'Dull Skin': ('skincare for dull skin', 'dull-skin'), 'Dry Lips': ('lip balm for dry lips', 'lip-balm-for-dry-lips'),
             'Dry Hair': ('hair mask for dry hair', 'hair-mask-for-dry-hair'), 'Body Hydration': ('body lotion for dry skin', 'body-lotion-for-dry-skin'), 'Anti-Dandruff': ('shampoo for dandruff', 'shampoo-for-dandruff'),
             'Frizzy Hair': ('hair serum for frizzy hair', 'hair-serum-for-frizzy-hair'), 'Hair Thinning & Loss': ('hair oil for hair loss', 'hair-oil-for-hair-loss'), 'Sun Protection': ('natural sunscreen', 'natural-sunscreen'),
             'Acne & Blemishes': ('face wash for oily skin', 'face-wash-for-oily-skin'), 'Oily Skin': ('face wash for oily skin', 'face-wash-for-oily-skin')}
        return m.get(con, ('natural skincare gift set', 'gifts-for-her'))
    return ('natural skincare', 'natural-skincare')
RISK = re.compile(r'pigment|dark spots?|wrinkle|acne scars?|regrow|re-grow|cure|treat(s|ment of)?\b|heal|clinical|guarantee|eliminat|prevent|within \d+ days|100% results|permanent|remove(s)? (dark|pigment|acne|scars)|acne-free|anti-?bacterial|fungal|hair fall within|stops? hair', re.I)
def facts(src):
    if not src: return [], False
    lines = [l.strip(' -•\t') for l in src['desc'].split('\n')]
    lines = [re.sub(r'^[A-Z][^:]{2,45}:\s*', '', l) for l in lines if len(l) > 25]
    risky = any(RISK.search(l) for l in lines)
    ok = [l for l in lines if not RISK.search(l) and not re.search(r'review|raves|wondering|whether you|deserves', l, re.I)]
    out = []
    for l in ok:
        l = l.rstrip(' .') + '.'; l = l[0].upper() + l[1:]
        if len(l) <= 170: out.append(l)
    return out[:2], risky
def tcase(s): return ' '.join(w if w in ('for', 'and', 'of') else w[:1].upper() + w[1:] for w in s.split())
def noun(kw):
    if kw.startswith(('skincare', 'natural skincare')): return 'natural ' + kw.replace('natural ', '')
    if kw.startswith('natural'): return 'a ' + kw
    return 'a natural ' + kw
wb = openpyxl.load_workbook(WB)
if 'VD Content' in wb.sheetnames: del wb['VD Content']
ws = wb.create_sheet('VD Content', 2)
A = 'Arial'; F = Font(name=A, size=10); B = Font(name=A, size=10, bold=True); H = Font(name=A, bold=True, color='FFFFFF', size=10); HF = PatternFill('solid', fgColor='1F3B34'); BR = PatternFill('solid', fgColor='9A7536')
cols = ['ID', 'Brand', 'Product', 'Type', 'Concern', 'Target keyword', 'Target vol.', 'Target KD', 'Brand keyword', 'Brand vol.', 'Collection page (link to)', 'H1', 'Meta title', 'Title len', 'Meta description', 'Desc len', 'Intro paragraph (first lines of product copy)', 'FAQ ideas (Google autocomplete)', 'Source description', 'Claims check']
widths = [9, 9, 34, 18, 18, 28, 9, 8, 10, 9, 30, 34, 52, 7, 70, 7, 90, 44, 14, 30]
for i, (c, w) in enumerate(zip(cols, widths), 1):
    x = ws.cell(row=1, column=i, value=c); x.font = H; x.fill = BR if c in ('Target vol.', 'Target KD', 'Brand vol.') else HF; x.alignment = Alignment(wrap_text=True, vertical='center')
    ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
ws.row_dimensions[1].height = 32; ws.freeze_panes = 'D2'
SD = "'Semrush Data'"
look = lambda c, col: f'=IFERROR(INDEX({SD}!${col}$2:${col}$5000,MATCH({c},{SD}!$A$2:$A$5000,0)),"")'
thin = Border(bottom=Side(style='thin', color='E3DACD'))
for r, p in enumerate(VD, 2):
    src = match(p); fx, risky = facts(src)
    kw, slug = target(p['Type'], p['Concern / scent'])
    text = ((src or {}).get('desc', '') + ' ' + p['Product name']).lower()
    oily, dry = bool(re.search(r'oily|acne|sebum|blemish', text)), bool(re.search(r'\bdry\b|dehydrat', text))
    if kw == 'moisturiser for dry skin' and oily and not dry: kw, slug = 'moisturiser for oily skin', 'moisturiser'
    if kw == 'face wash for dry skin' and oily and not dry: kw, slug = 'face wash for oily skin', 'face-wash-for-oily-skin'
    if kw == 'face wash for oily skin' and dry and not oily: kw, slug = 'face wash for dry skin', 'face-wash-for-dry-skin'
    h1 = p['Draft H1']; brand = p['Brand']; KW = tcase(kw)
    tc = [f'{h1} | {KW}', f'{h1} – {KW} | Swansura', f'{h1} | {brand} UK | Swansura', f'{h1} | Swansura UK']
    title = next((t for t in tc[:2][::-1] if 50 <= len(t) <= 60), None) or next((t for t in tc if len(t) <= 60), h1[:47] + ' | Swansura')
    short = fx[0].rstrip('.') if fx else ''
    leads = [f'Buy {h1}, {noun(kw)} from {brand}.', f'{h1}: {noun(kw)} from {brand}.']
    words = short.split(); s2 = ' '.join(words[:12]).rstrip(',;:') if short else ''
    mids = [f' {short}.' if short and len(short) < 95 else '', f' {s2}.' if s2 else '', '']
    tails = [f' Official UK partner: 100% authentic, shipped from the UK, free delivery over £50.', f' Official UK partner, 100% authentic, free UK delivery over £50.', ' 100% authentic, shipped from the UK, free delivery over £50.', ' Free UK delivery over £50.']
    combos = [re.sub(r'\s+', ' ', a + b + c).strip() for a in leads for b in mids for c in tails]
    ok = [x for x in combos if 140 <= len(x) <= 160]
    desc = max(ok, key=len) if ok else min(combos, key=lambda x: abs(len(x) - 150) if len(x) <= 160 else 999)
    intro = f'{h1} is {noun(kw)} from {brand}. ' + ' '.join(fx) + f' Swansura is an official UK partner of {brand}, so every order is 100% authentic and ships from the UK.'
    intro = re.sub(r'\s+', ' ', intro)
    row = [p['ID'], brand, p['Product name'], p['Type'], p['Concern / scent'], kw, None, None, brand.lower(), None, f'/collections/{slug}', h1, title, None, desc, None, intro, p['Google autocomplete phrases'] or '', ('Matched' if fx else 'Matched – no usable facts, add 1–2 from pack') if src else 'Not found – write from pack', 'Review: source copy has medical-style claims (kept out of intro)' if risky else '']
    for c, v in enumerate(row, 1):
        x = ws.cell(row=r, column=c, value=v); x.font = F; x.border = thin; x.alignment = Alignment(wrap_text=c in (15, 17, 18, 20), vertical='top')
    ws.cell(row=r, column=7, value=look(f'F{r}', 'C')).number_format = '#,##0'
    ws.cell(row=r, column=8, value=look(f'F{r}', 'D'))
    ws.cell(row=r, column=10, value=look(f'I{r}', 'C')).number_format = '#,##0'
    ws.cell(row=r, column=14, value=f'=LEN(M{r})'); ws.cell(row=r, column=16, value=f'=LEN(O{r})')
    for c in (7, 8, 10, 14, 16): ws.cell(row=r, column=c).font = F
    ws.cell(row=r, column=6).font = B
n = len(VD) + 1
red = Font(name=A, size=10, color='C00000', bold=True)
ws.conditional_formatting.add(f'N2:N{n}', FormulaRule(formula=['OR(N2<50,N2>60)'], font=red))
ws.conditional_formatting.add(f'P2:P{n}', FormulaRule(formula=['OR(P2<140,P2>160)'], font=red))
ws.conditional_formatting.add(f'T2:T{n}', FormulaRule(formula=['T2<>""'], fill=PatternFill('solid', fgColor='FCE4D6')))
ws.conditional_formatting.add(f'S2:S{n}', FormulaRule(formula=['S2<>"Matched"'], fill=PatternFill('solid', fgColor='FCE4D6')))
ws.auto_filter.ref = f'A1:T{n}'
wb.save(WB); print('rows', len(VD), 'matched', sum(1 for p in VD if match(p)))
